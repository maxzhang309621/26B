"""Reorder Q4 figures (image then interpretation) and insert 8.3 drill paths."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

SRC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订3.docx"
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订4.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题4(1).docx"
)
ALT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订.docx"
FIG = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\figures")
FIG16 = FIG / "p4-20260912-210603_path.png"
FIG12 = FIG / "p4-20260912-210518_path.png"
FIG10 = FIG / "p4-20260912-205941_path.png"
DUMP = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q4_figorder.txt")
WIDTH = 13.6


def blank_after(anchor, style_para: Paragraph) -> Paragraph:
    new_elm = deepcopy(style_para._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    anchor._p.addnext(new_elm)
    return Paragraph(new_elm, style_para._parent)


def strip_omath(para: Paragraph) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def set_plain(para: Paragraph, text: str) -> None:
    strip_omath(para)
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def find_start(doc: Document, prefix: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def drop(para: Paragraph) -> None:
    parent = para._p.getparent()
    if parent is not None:
        parent.remove(para._p)


def move_after(anchor: Paragraph, *items: Paragraph) -> None:
    for para in reversed(items):
        anchor._p.addnext(para._p)


def insert_figure(anchor, body: Paragraph, cap_style: Paragraph, png: Path,
                  caption: str, note: str, width_cm: float = WIDTH) -> Paragraph:
    note_p = blank_after(anchor, body)
    cap_p = blank_after(anchor, cap_style)
    img_p = blank_after(anchor, body)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(str(png), width=Cm(width_cm))
    cap_p.add_run(caption)
    note_p.add_run(note)
    return note_p


def main() -> None:
    for png in (FIG16, FIG12, FIG10):
        if not png.exists():
            raise FileNotFoundError(png)

    copy2(SRC, OUT)
    doc = Document(str(OUT))
    body = find_start(doc, "求解对应默认入口")

    h821 = find_start(doc, "8.2.1")
    cap3 = find_start(doc, "图 3  ")
    img3 = Paragraph(cap3._p.getprevious(), cap3._parent)
    note3 = find_start(doc, "左图虚线为任务圆")
    cap3c = find_start(doc, "图 3（续）")
    note3c = find_start(doc, "柱长为行驶与驻留之和")
    rank = find_start(doc, "在该口径下：7 内环")

    cap4 = find_start(doc, "图 4  ")
    img4 = Paragraph(cap4._p.getprevious(), cap4._parent)
    note4 = find_start(doc, "左上：等圆变为盘")
    vis = find_start(doc, "P ∈ D(G")

    p60 = find_start(doc, "覆盖—服务两阶段沿用问题三")
    cap5 = find_start(doc, "图 5  ")
    img5 = Paragraph(cap5._p.getprevious(), cap5._parent)
    note5 = find_start(doc, "覆盖四段结束后进入锁定开放路")

    p76 = find_start(doc, "覆盖期内的路径改动只作用")
    cap6 = find_start(doc, "图 6  ")
    img6 = Paragraph(cap6._p.getprevious(), cap6._parent)
    note6 = find_start(doc, "左图虚线圈出")

    p91 = find_start(doc, "代表局表明")
    p_inv = find_start(doc, "反演验证用官方演练日志")
    cap_v7 = find_start(doc, "图 7  ")
    img_v7 = Paragraph(cap_v7._p.getprevious(), cap_v7._parent)
    note_v7 = find_start(doc, "稀源局覆盖段更完整")
    cap_v8 = find_start(doc, "图 8  ")
    img_v8 = Paragraph(cap_v8._p.getprevious(), cap_v8._parent)
    note_v8 = find_start(doc, "密源局在覆盖后进入局部簇清除")

    move_after(h821, img3, cap3, note3)
    move_after(note3c, rank)
    move_after(vis, img4, cap4, note4)
    move_after(p60, img5, cap5, note5)
    move_after(p76, img6, cap6, note6)

    set_plain(
        p91,
        "表 1 按源数只列少数、多数、最多三类。少数取本批最稀的十源，7400 s / 740 s·源，固定骨干分摊抬高单位耗时；"
        "多数取出现四次的十二源中与表 1 相同的一局，7583 s / 632 s·源；"
        "最多取十六源上界中最快的一局，6201 s / 388 s·源。"
        "三局航迹见图 7、图 7（续）与图 8，均来自 21:06 批次官方演练日志，与表 1 为同一局；"
        "源位由清除点与示向反演还原，图中全向/定向符号是日志推断，官方组成以表注为准。",
    )
    insert_figure(
        p91, body, cap6, FIG10,
        "图 8  演练反演：十源少数局（表 1 少数）",
        "十源局外环听点被走得更完整，覆盖骨干在源数上分摊不足，秒/源最高。"
        "官方组成为 7 全向 + 3 定向；清除成功点仍落在听点环与途听邻域，说明稀源并不缩短必须走完的覆盖骨架。",
    )
    insert_figure(
        p91, body, cap6, FIG12,
        "图 7（续）  演练反演：十二源多数局（表 1 多数）",
        "该局官方为 12 定向、0 全向，秒/源 632，介于十源与十六源之间。"
        "外环采样被真正消耗，定向源清除折返出现在覆盖结束之后的局部簇，与先搜后清相位一致。",
    )
    insert_figure(
        p91, body, cap6, FIG16,
        "图 7  演练反演：十六源最快局（表 1 最多）",
        "十六源局总时 6201 s，秒/源 388，为本批最低。外环未主导全程，清除段嵌入覆盖之后的近场簇，"
        "途清与无跳点局部开放路更容易触发。新增源主要提高骨干利用率，而不是按比例拉长总时。",
    )

    set_plain(
        p_inv,
        "反演验证用官方演练日志回放路径与清除事件，核对策略是否按宣称相位执行。"
        "图 7 至图 8 三局均来自 21:06 批次：覆盖搜索段沿 7+12 听点与出发途听展开，"
        "定位清除段主要出现在覆盖结束之后；清除成功点落在估计服务点邻域，且每局清除数等于官方 jammer_count。"
        "路径形态已在 8.3 节给出，此处不再重复插图。",
    )
    for para in (img_v7, cap_v7, note_v7, img_v8, cap_v8, note_v8):
        drop(para)

    saved = OUT
    try:
        doc.save(str(OUT))
    except PermissionError:
        saved = Path(str(OUT).replace("修订4", "修订4b"))
        doc.save(str(saved))
    copies = {"saved": saved.name}
    for dest, key in ((WX, "wechat"), (ALT, "修订"), (SRC, "修订3")):
        try:
            copy2(saved, dest)
            copies[key] = True
        except OSError:
            copies[key] = False

    d2 = Document(str(saved))
    lines = [f"paras={len(d2.paragraphs)} copies={copies}"]
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        n_img = len(para._p.findall(".//" + qn("w:drawing")))
        extra = f" [img={n_img}]" if n_img else ""
        if t.startswith((
            "8.2.", "图 ", "左图", "左上", "覆盖四段", "柱长", "在该口径下",
            "P ∈", "覆盖—服务", "本问只补充", "覆盖期内的路径", "若直接对全部",
            "代表局表明", "十源局", "该局官方", "十六源局", "反演验证用",
            "8.3", "8.4.2", "近心朝外",
        )) or extra:
            lines.append(f"{i:04d}{extra} {t[:160]}")
    DUMP.write_text("\n".join(lines), encoding="utf-8")
    print("saved", saved.name)
    print("copies", copies)


if __name__ == "__main__":
    main()
