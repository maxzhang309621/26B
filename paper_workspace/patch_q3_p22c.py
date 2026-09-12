from pathlib import Path
from shutil import copy2

from docx import Document

DOC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订3.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(2).docx"
)

doc = Document(str(DOC))
p = doc.paragraphs[22]
p.runs[12].text = (
    "，不随源的空间分布自适应调整。"
    "图 2 左仅给出本文正六边形检测点，右图为最不利覆盖的几何核对；"
    "图 2（续）"
)
p.runs[13].text = ""
p.runs[14].text = "用排序柱状图给出耗时对照，不再列表、也不并列八边形分布图。"
doc.save(str(DOC))
try:
    copy2(DOC, WX)
    copied = True
except OSError:
    copied = False
d2 = Document(str(DOC))
Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q3_p22c.txt").write_text(
    d2.paragraphs[22].text + "\n\n" + d2.paragraphs[24].text + "\n\n" + d2.paragraphs[30].text,
    encoding="utf-8",
)
print("copied", copied, "tables", len(d2.tables))
