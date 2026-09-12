"""Clear leftover ' 2' and add spacing in the time-proof paragraph."""
from pathlib import Path
from shutil import copy2

from docx import Document

DOC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace\2026B-问题3(2)-修订.docx")
SRC = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(2).docx"
)

doc = Document(str(DOC))
p = doc.paragraphs
# original para 22, still index 22
# original para 22: keep "表 2", drop the leftover second "2"
if len(p[22].runs) > 15:
    p[22].runs[15].text = ""
# inserted time-proof paragraph is 23
if p[23].runs:
    p[23].runs[0].text = p[23].runs[0].text.replace("故T(n,ρ)=", "故 T(n,ρ) = ")
doc.save(str(DOC))
try:
    copy2(DOC, SRC)
    copied = True
except OSError:
    copied = False
d2 = Document(str(DOC))
Path = Path  # noqa
from pathlib import Path as P
P(r"D:\maxzhang\python files\Math-modeling\26B\output\_q3v2_p22.txt").write_text(
    d2.paragraphs[22].text + "\n\n" + d2.paragraphs[23].text,
    encoding="utf-8",
)
print("copied", copied)
