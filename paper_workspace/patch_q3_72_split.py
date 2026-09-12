"""Restructure Q3 7.2: hexagon → search-then-clear → edge / enroute / optical."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

SRC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订4.docx"
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订5.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(2).docx"
)
ALT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订.docx"
FIG4 = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\figures\q3_edge_enroute_algo.png")
DUMP = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q3_72_new.txt")


def blank_after(anchor, style_para: Paragraph) -> Paragraph:
    new_elm = deepcopy(style_para._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    if hasattr(anchor, "_p"):
        anchor._p.addnext(new_elm)
    else:
        anchor.addnext(new_elm)
    return Paragraph(new_elm, style_para._parent)


def strip_omath(para: Paragraph) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def replace_picture(para: Paragraph, png: Path, width_cm: float) -> None:
    p_pr = para._p.find(qn("w:pPr"))
    for child in list(para._p):
        if child is not p_pr:
            para._p.remove(child)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(png), width=Cm(width_cm))


def find_by_prefix(doc: Document, prefix: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def main() -> None:
    copy2(SRC, OUT)
    doc = Document(str(OUT))

    intro = find_by_prefix(doc, "鉴于源位置先验未知")
    intro.runs[0].text = (
        "鉴于源位置先验未知、环境以协议黑盒形式给出、且听点数目极少，本问采用构造性求解而非开环轨迹优化。"
        "求解顺序为：先由等圆覆盖确定正六边形检测点；再规定先搜后清的覆盖—服务相位，并在覆盖结束后以滚动开放路实施批量清除；"
        "然后分述环边补测、顺路清除与光学网格各自要解决的问题、门控公式与示意图。"
        "路径组织方式的演进与淘汰依据见本节末。不以整数规划、遗传算法或强化学习作为主求解器。"
    )

    h222 = find_by_prefix(doc, "7.2.2")
    h222.runs[1].text = "先搜后清"

    p32 = find_by_prefix(doc, "求解过程首先进入任务区")
    strip_omath(p32)
    p32.runs[0].text = (
        "求解过程首先进入任务区并读取局内源数；于原点对全部二十个信道各实施一次检测，遇近场指示则立即清除。"
        "随后依正六边形听点次序巡游，每点仅扫描尚未排除的信道：一旦闻及信号，即记入待清集合，覆盖期内不实施专程折返清除。"
        "覆盖骨架走完或已听数目达到局内源数后，再转入批量清除。图 3 对照“闻及即清”的星形折返与先搜后清的先环后批结构；"
        "环边补测、顺路清除与光学网格是骨架上的三项有条件机制，分见后文。"
    )
    for run in p32.runs[1:]:
        run.text = ""

    p34 = find_by_prefix(doc, "则允许顺路清除")
    p34.runs[0].text = (
        "图 3 左在每次发现后即脱离覆盖骨干，增量路程逐次累积；右图先完成听点遍历，再统一实施清除。"
    )
    for run in p34.runs[1:]:
        run.text = ""

    note3 = find_by_prefix(doc, "左图在每次发现后即脱离覆盖骨干")
    batch = find_by_prefix(doc, "当听点遍历完毕或已听数目达到局内源数后")
    note3._p.addnext(batch._p)

    h_style = h222
    h223 = blank_after(batch, h_style)
    h223.add_run("7.2.3  ").bold = False
    # copy heading style by using same pPr; write full heading in one run
    h223.runs[0].text = "7.2.3  环边补测：把第二次示向摊进已有环边"
    edge = blank_after(h223, p32)
    edge.add_run(
        "环边补测要解决的问题是：若覆盖结束时大量信道仍只有一次示向，则服务点尚未成为可清除对象，"
        "扫完后再优化访问次序的空间很小，折返会被推迟到批量阶段集中爆发。"
        "设当前环边为 A→B。在边上取离散采样"
    )
    edge_f = blank_after(edge, p32)
    edge_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    edge_f.add_run("P(t)=(1−t)A+tB，t∈{0.25, 0.40, 0.50, 0.60, 0.75}。")
    edge2 = blank_after(edge_f, p32)
    edge2.add_run(
        "对恰有一次示向 (s₁, θ) 的信道，若后续尚未访问的环点均不落在问题二候选带 C(s₁,θ) 内，"
        "则在满足 P(t)∈C(s₁,θ) 的样本中选取横向偏离最大者作为补测点："
    )
    edge_f2 = blank_after(edge2, p32)
    edge_f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    edge_f2.add_run("t* = arg max | (P(t)−s₁) × u(θ) |。")
    edge3 = blank_after(edge_f2, p32)
    edge3.add_run(
        "这不是把移动轨迹当作连续监听，而是只在若干真实检测位置上增加一次合法第二站。"
        "后续顶点已能提供合法二测则不插。图 4 左给出采样点与候选带的几何关系。"
    )

    cap4 = find_by_prefix(doc, "图 4")
    cap4.runs[0].text = "图 4  环边补测的边上采样（左）与顺路清除的增量路程门控（右）"
    for run in cap4.runs[1:]:
        run.text = ""
    img4 = Paragraph(cap4._p.getprevious(), cap4._parent)
    replace_picture(img4, FIG4, 16.0)

    note4 = find_by_prefix(doc, "实线表示覆盖边")
    if note4.runs:
        note4.runs[0].text = (
            "左图：仅当恰有一次示向且后续顶点不能提供合法第二站时，在问题二候选带与环边交点上补测。"
            "右图：已有不少于两次示向时，仅当增量路程不超过"
        )
        if len(note4.runs) >= 4:
            note4.runs[-1].text = "且非近共线，才允许暂时离开环边 P→Q 前往服务点 C。"

    h224 = blank_after(note4, h_style)
    h224.add_run("7.2.4  顺路清除：用增量路程限制覆盖期内的有条件服务")
    enr = blank_after(h224, p32)
    enr.add_run(
        "顺路清除要解决的问题是：部分信道在环上已具备不少于两次示向、定位区已可清，"
        "若一律拖到覆盖结束再专程折返，会把本可就近完成的服务堆进批量阶段。"
        "记当前位置为 P、下一听点为 Q、估计服务点为 C。增量路程满足"
    )
    eq4 = find_by_prefix(doc, "(4)")
    enr._p.addnext(eq4._p)
    enr_after = blank_after(eq4, p32)
    enr_after.add_run(
        "则允许暂时离开环边执行清除，否则延至覆盖结束之后。"
        "该门控把可清除机会融入已有环边，避免在覆盖骨架上反复叠加专程折返。图 4 右给出 ΔL 的几何含义。"
    )

    optical = find_by_prefix(doc, "当定位区尺度已充分缩小")
    prev = Paragraph(optical._p.getprevious(), optical._parent)
    h225 = blank_after(prev, h_style)
    h225.add_run("7.2.5  光学网格：仅在小定位区上作有界试清")
    # p42 already follows; add half-diagonal to the note
    note5 = find_by_prefix(doc, "网格仅服务于")
    note5.runs[0].text = (
        "网格仅服务于“定位区已很小但仍略超二十米”的情形。格距取 25 m，正方形格心至角点距离为 "
        "25√2/2≈17.68 m，小于 20 m 清除半径，故格心试清成功即覆盖该格；最多 8 格、试 1 次，miss 即停。"
        "对大张角、近共线或跨度过大的区域不启用铺网，改行补测。清除充分条件始终比较最小包围圆半径与 20 m，"
        "不以直径 40 m 代替。"
    )

    mech = find_by_prefix(doc, "上述结构的机理在于")
    for run in mech.runs:
        if run.text and "7.2.3" in run.text:
            run.text = run.text.replace("7.2.3", "7.2.6")

    evo = find_by_prefix(doc, "7.2.3")
    # after inserting 7.2.3 环边补测, the evolution heading still starts with 7.2.3  从边扫边清
    if "从边扫边清" in evo.text:
        evo.runs[0].text = "7.2.6  "
        if len(evo.runs) > 1:
            evo.runs[1].text = "从边扫边清到先搜后清的演进与取舍"
        else:
            evo.runs[0].text = "7.2.6  从边扫边清到先搜后清的演进与取舍"
    else:
        for para in doc.paragraphs:
            if para.text.strip().startswith("7.2.3") and "从边扫边清" in para.text:
                para.runs[0].text = "7.2.6  "
                if len(para.runs) > 1:
                    para.runs[1].text = "从边扫边清到先搜后清的演进与取舍"

    doc.save(str(OUT))
    copies = {}
    for dest, key in ((WX, "wechat"), (ALT, "修订")):
        try:
            copy2(OUT, dest)
            copies[key] = True
        except OSError:
            copies[key] = False

    d2 = Document(str(OUT))
    lines = [f"paras={len(d2.paragraphs)} copies={copies}"]
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        if t.startswith(("7.2", "图 3", "图 4", "图 5", "求解过程", "环边补测", "顺路清除", "网格仅", "当听点", "当定位", "上述结构", "P(t)", "t*")):
            lines.append(f"{i:04d} {t[:180]}")
    DUMP.write_text("\n".join(lines), encoding="utf-8")
    print("saved", OUT)
    print("copies", copies)


if __name__ == "__main__":
    main()
