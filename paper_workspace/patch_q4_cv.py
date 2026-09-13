"""Move Q4 listen-point time experiments from 8.2.1 into 8.4.1."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订4.docx"
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订5.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题4(1).docx"
)
ALT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订.docx"
DUMP = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q4_cv.txt")


def blank_after(anchor, style_para: Paragraph) -> Paragraph:
    new_elm = deepcopy(style_para._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    anchor._p.addnext(new_elm)
    return Paragraph(new_elm, style_para._parent)


def strip_omath(para: Paragraph) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def set_plain(para: Paragraph, text: str) -> None:
    strip_omath(para)
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def find_start(doc: Document, prefix: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def move_after(anchor: Paragraph, *items: Paragraph) -> None:
    for para in reversed(items):
        anchor._p.addnext(para._p)


def main() -> None:
    copy2(SRC, OUT)
    doc = Document(str(OUT))

    intro = find_start(doc, "求解对应默认入口")
    h821 = find_start(doc, "8.2.1")
    p_close = find_start(doc, "综合几何门槛与耗时排序")
    p_time = find_start(doc, "覆盖巡游的虚拟时间须把几何路程")
    p_tform = find_start(doc, "T = L/5")
    p_ldef = find_start(doc, "其中 L 为原点经内环再到外环")
    cap3c = find_start(doc, "图 3（续）")
    img3c = Paragraph(cap3c._p.getprevious(), cap3c._parent)
    note3c = find_start(doc, "柱长为行驶与驻留之和")
    rank = find_start(doc, "在该口径下：7 内环")
    h841 = find_start(doc, "8.4.1")
    p_dense = find_start(doc, "交叉验证回答：同一几何目标下")
    p_strat = find_start(doc, "策略交叉上")
    body = intro

    set_plain(
        intro,
        "求解对应默认入口 hexbatch。源位置与定向朝向均先验未知，环境以协议黑盒给出，"
        "故采用构造性求解：先由覆盖几何确定双环检测点，再相对问题三改写定向可听条件与服务门控；"
        "随后沿用先搜后清相位，展开定向前瓣校正与无跳点开放路动态规划。"
        "不以整数规划或开环轨迹优化作为主求解器。多种检测点分布的覆盖巡游耗时对照见 8.4.1 节。",
    )
    set_plain(h821, "8.2.1  双环检测点：覆盖几何与听点构造")
    set_plain(
        p_close,
        "综合上述覆盖几何门槛，听点构造写为原点、内正七边形左端半径与外正十二边形；途听规则如下。"
        "覆盖巡游时间由路程与停靠驻留共同决定，多种分布的耗时对照见 8.4.1 节。",
    )

    lead = blank_after(h841, body)
    lead.add_run(
        "交叉验证分三层。第一层固定覆盖口径，比较多种检测点分布的巡游下界，"
        "检验 8.2.1 节由几何门槛锁住的 7+12 同径向点集是否仍是可行方案中的最短。"
        "第二层核对该点集的密采样完备性。第三层换航路、换相位或关掉校正/顺路清，"
        "看全清结论是否仍一致。覆盖下界不含定位清除折返，与 8.3 节官方演练成绩分开书写。"
    )

    set_plain(
        p_time,
        "检测点分布交叉把几何路程与协议动作一并计入。原点对 1–20 信道全扫耗时 119 s；"
        "其后每个停靠在最坏情形下仍须扫描全部未知信道，单点驻留 120 s。"
        "以 5 m/s 折合行驶时间，停靠数为 N 的先搜后清开放巡游满足",
    )
    set_plain(cap3c, "图 9  双环覆盖巡游虚拟时间（最坏满扫 20 信道）")
    set_plain(
        rank,
        "同一口径下交叉比较多种检测点分布。7 内环取覆盖区间左端、12 外环取 1865 m 且同径向时，"
        "行驶 3655 s、驻留 2399 s，合计 6054 s，是前瓣核验通过方案中的最短。"
        "交错相位同点集行驶增至 3707 s，合计 6106 s。"
        "内环增至正八、正九或外环增至十四边形，停靠驻留立刻把总时抬到 6180–6328 s。"
        "若把 900 m 计入听点集以使正六内环通过前瓣，停靠增至 25、总时 7888 s；"
        "旧方案 8×1200+12×2100+途听环为 8773 s。"
        "7+11×1900 合计 5958 s 更短，但前瓣核验未通过，不能作为可行最短。"
        "图 9 按总时排序给出上述对照。因此交叉验证支持把完备性听点锁在 7+12 同径向："
        "更短的 11 点外环前瓣不过，更密的多边形驻留抬时，途听环补前瓣则把总时抬到 7888 s。",
    )

    p_dense.runs[0].text = "完备性口径交叉上，全向侧原点与内七边形在 "
    p_strat.runs[3].text = " 的差距可由覆盖巡游下界（6054 s）解释，而不是探测公式更换。"

    move_after(lead, p_time, p_tform, p_ldef, img3c, cap3c, note3c, rank)

    saved = OUT
    try:
        doc.save(str(OUT))
    except PermissionError:
        saved = Path(str(OUT).replace("修订5", "修订5b"))
        doc.save(str(saved))
    copies = {"saved": saved.name}
    for dest, key in ((WX, "wechat"), (ALT, "修订")):
        try:
            copy2(saved, dest)
            copies[key] = True
        except OSError:
            copies[key] = False

    d2 = Document(str(saved))
    lines = [f"paras={len(d2.paragraphs)} copies={copies}"]
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        n_img = len(para._p.findall(".//" + qn("w:drawing")))
        extra = f" [img={n_img}]" if n_img else ""
        if t.startswith((
            "8.2.1", "8.2.2", "求解对应", "综合上述", "覆盖巡游路程",
            "定向源把可听", "访问顺序", "设内顶点", "8.4.1", "交叉验证分三层",
            "检测点分布交叉", "T =", "其中 L", "图 9", "柱长", "同一口径下",
            "完备性口径", "策略交叉", "8.4.2", "图 3",
        )) or extra:
            if 18 <= i <= 50 or i >= 100 or extra:
                lines.append(f"{i:04d}{extra} {t[:170]}")
    DUMP.write_text("\n".join(lines), encoding="utf-8")
    print("saved", saved.name)
    print("copies", copies)


if __name__ == "__main__":
    main()
