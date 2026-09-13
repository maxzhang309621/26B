"""Second-pass wording fixes on the already-revised Q3 docx."""
from pathlib import Path
from shutil import copy2

from docx import Document

REV = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace\2026B-问题3-修订.docx")
ORIG = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(1).docx"
)
CHECK = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\q3_paper_check.txt")


def main() -> None:
    doc = Document(str(REV))
    p = doc.paragraphs

    p[10].runs[10].text = " 为缺省），或覆盖航路走尽。三阶段示意如图"
    p[56].runs[3].text = ""
    p[56].runs[25].text = "访问顺序的"
    p[72].runs[5].text = "边补二测（20260912）"

    for para in p:
        if para.text.startswith("综上"):
            for run in para.runs:
                if "清除失败" in run.text:
                    run.text = run.text.replace(" 清除失败 ", "清除失败")
            break

    cell = doc.tables[3].rows[-1].cells[-1]
    for para in cell.paragraphs:
        if para.runs:
            para.runs[0].text = "50（合计）"
            for run in para.runs[1:]:
                run.text = ""
        elif "—" in para.text or para.text.strip() in {"", "—"}:
            para.add_run("50（合计）")

    doc.save(str(REV))
    try:
        copy2(REV, ORIG)
        copied = True
    except OSError:
        copied = False

    doc = Document(str(REV))
    lines = [
        doc.paragraphs[10].text[-50:],
        doc.paragraphs[56].text[:40],
        "拧:" + ("访问顺序" if "访问顺序" in doc.paragraphs[56].text else "FAIL"),
        doc.paragraphs[72].text[doc.paragraphs[72].text.find("证据") : doc.paragraphs[72].text.find("证据") + 30],
        "table3=" + doc.tables[3].rows[-1].cells[-1].text,
        "copied=" + str(copied),
        "omath=" + str(sum(x._p.xml.count("oMath") for x in doc.paragraphs)),
    ]
    CHECK.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
