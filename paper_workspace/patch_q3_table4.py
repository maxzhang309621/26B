"""Shrink 表4 to three representative jammer-count cases."""
from __future__ import annotations

from pathlib import Path
from shutil import copy2

from docx import Document

SRC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订6.docx"
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订7.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(2).docx"
)
ALT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订.docx"


def set_cell(cell, text: str) -> None:
    paras = cell.paragraphs
    if not paras:
        return
    if paras[0].runs:
        paras[0].runs[0].text = text
        for run in paras[0].runs[1:]:
            run.text = ""
    else:
        paras[0].add_run(text)
    for para in paras[1:]:
        for run in para.runs:
            run.text = ""


def strip_omath(para) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def find_start(doc: Document, prefix: str):
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def main() -> None:
    copy2(SRC, OUT)
    doc = Document(str(OUT))

    cap = find_start(doc, "表 4")
    cap.runs[0].text = "表 4  官方演练代表性源数情形（少数 / 多数 / 最多）"
    for run in cap.runs[1:]:
        run.text = ""

    p79 = find_start(doc, "主结果取官方问题三演练")
    for run in p79.runs:
        if run.text and "逐局结果见表" in run.text:
            run.text = run.text.replace("逐局结果见表", "代表性源数情形见表")

    tbl = doc.tables[2]
    while len(tbl.rows) > 4:
        tbl._tbl.remove(tbl.rows[-1]._tr)
    headers = ["情形", "源数", "虚拟时间/s", "秒/源", "说明"]
    rows = [
        ["少数", "11", "3806", "346", "本批最稀，骨干摊薄不足"],
        ["多数", "13", "3855", "297", "本批最常见（5/15 局），近全样本均值"],
        ["最多", "16", "3866", "242", "源数上界，本批最优单局"],
    ]
    for j, h in enumerate(headers):
        set_cell(tbl.rows[0].cells[j], h)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            set_cell(tbl.rows[i].cells[j], val)

    p82 = find_start(doc, "由表")
    if p82.runs:
        p82.runs[0].text = "结合十五局全清与表"

    p83 = find_start(doc, "在源数取")
    strip_omath(p83)
    p83.runs[0].text = (
        "表 4 按源数只列三类代表局。少数取本批最稀的十一源，秒/源 346、总时约 3806 s，"
        "六边形骨干成本几乎不随源数下降。多数取出现次数最多的十三源（5/15 局），代表局秒/源 297，"
        "接近全样本 295 s/源。最多取十六源上界，代表局秒/源 242、总时约 3866 s，与十一源局总时接近，"
        "表明新增源主要提高骨干利用率，而非按比例拉长总时。本地十六源扇区聚集路径如图 8："
        "听点环仍被完整遍历，清除折返集中于源密集扇区。图 8 仅作形态说明，不代替表 4 的演练口径。"
    )
    for run in p83.runs[1:]:
        run.text = ""

    p87 = find_start(doc, "偏稀情形下")
    strip_omath(p87)
    p87.runs[0].text = (
        "少数局的路径形态见图 9：本地十源周向分散、场内大片留白时，覆盖骨干依旧完整，"
        "批量阶段跨空白区的边长增加，秒/源与官方十一源局同处偏高量级。"
        "稀疏布局主要恶化单位源耗时，而非削弱全清能力；不宜依据源数动态删减听点，以免破坏最不利覆盖保证。"
    )
    for run in p87.runs[1:]:
        run.text = ""

    doc.save(str(OUT))
    copies = {}
    for dest, key in ((WX, "wechat"), (ALT, "修订")):
        try:
            copy2(OUT, dest)
            copies[key] = True
        except OSError:
            copies[key] = False
    d2 = Document(str(OUT))
    t = d2.tables[2]
    print("saved", OUT.name, "copies", copies, "table_rows", len(t.rows))
    for row in t.rows:
        print(" | ".join(c.text.replace("\n", " ") for c in row.cells))


if __name__ == "__main__":
    main()
