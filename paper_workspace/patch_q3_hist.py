"""Replace 覆盖证书 wording, drop table 2, insert ranked time chart."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

SRC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订.docx"
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订3.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(2).docx"
)
PNG = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\figures\q3_ngon_time_ranked.png")
EXTRACT = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q3_after_hist.txt")

REPLACES = (
    ("覆盖证书路程", "覆盖骨架路程"),
    ("证书路程", "覆盖骨架路程"),
    ("覆盖证书", "覆盖完备性判据"),
    ("证书成立", "完备性成立"),
    ("证书通过", "判据成立"),
)


def apply_replaces(text: str) -> str:
    for old, new in REPLACES:
        text = text.replace(old, new)
    return text


def rewrite_runs(doc: Document) -> None:
    for para in doc.paragraphs:
        for run in para.runs:
            if run.text:
                run.text = apply_replaces(run.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.text:
                            run.text = apply_replaces(run.text)


def insert_after(after_para, style_para=None) -> Paragraph:
    src = style_para or after_para
    new_elm = deepcopy(src._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    after_para._p.addnext(new_elm)
    return Paragraph(new_elm, after_para._parent)


def main() -> None:
    copy2(SRC, OUT)
    doc = Document(str(OUT))
    rewrite_runs(doc)

    cap = None
    for para in doc.paragraphs:
        if para.text.startswith("表 2"):
            cap = para
            break
    if cap is None:
        raise RuntimeError("table 2 caption not found")
    if cap.runs:
        cap.runs[0].text = (
            "图 2（续）  可行正 n 边形覆盖巡游耗时（最小可行半径，按总时升序）"
        )
        for run in cap.runs[1:]:
            run.text = ""

    # delete the n-gon comparison table (second table)
    tbl = doc.tables[1]._tbl
    tbl.getparent().remove(tbl)

    img_p = insert_after(cap, style_para=cap)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_p.add_run()
    run.add_picture(str(PNG), width=Cm(15.2))

    note = insert_after(img_p, style_para=cap)
    note.add_run(
        "各环取刚能覆盖的最小半径。蓝段为行驶，绿段为原点全扫与环点信道驻留；"
        "正六边形总时最短。本文采用 1150 m，总时约 2219 s，仍低于其余可行环。"
        "五点环越界，不进入比较。"
    )

    for para in doc.paragraphs:
        t = para.text
        if "图 2 左仅给出" in t:
            for run in para.runs:
                run.text = run.text.replace("表 2 ", "图 2（续）")
                run.text = run.text.replace(
                    "按几何覆盖与动作驻留合计给出耗时对照",
                    "用排序柱状图给出耗时对照",
                )
        if para.text.startswith("图 2  "):
            if para.runs:
                para.runs[0].text = "图 2  正六边形检测点几何分布（左）与最不利覆盖的几何核对（右）"
                for run in para.runs[1:]:
                    run.text = ""
        if "右图为本地" in para.text or "右图为最不利" in para.text:
            if para.runs:
                para.runs[0].text = (
                    "左图为本文采用的原点加正六边形 6×1150 m（相位 10°）检测点；浅盘为半径 1000 m 可听范围。"
                    "右图为最不利覆盖的几何核对：边界点到最近环点 988.5 m，余量约 11.5 m，完备性成立。"
                    "耗时比较见图 2（续），不再列表。"
                )
                for run in para.runs[1:]:
                    run.text = ""
        if para.text.startswith("图 8"):
            if para.runs:
                para.runs[0].text = "图 8  最不利覆盖的几何核对：最不利边界点落入可听盘内"
                for run in para.runs[1:]:
                    run.text = ""

    doc.save(str(OUT))
    lines = []
    d2 = Document(str(OUT))
    lines.append(f"paragraphs={len(d2.paragraphs)} tables={len(d2.tables)}")
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        if any(k in t for k in ("证书", "完备", "图 2", "表 2", "7.2.1", "图 8")):
            lines.append(f"{i:04d} {t[:180]}")
    EXTRACT.write_text("\n".join(lines), encoding="utf-8")
    try:
        copy2(OUT, WX)
        copied = True
    except OSError:
        copied = False
    print("saved", OUT)
    print("copied_wx", copied)


if __name__ == "__main__":
    main()
