from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

doc = Document(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace\2026B-问题4(1)-修订3.docx")
lines = [f"paras={len(doc.paragraphs)} tables={len(doc.tables)}"]
for i, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    n_img = len(para._p.findall(".//" + qn("w:drawing")))
    extra = f" [img={n_img}]" if n_img else ""
    if t or extra:
        lines.append(f"{i:04d}{extra} {t[:220]}")
Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q4_rev3_full.txt").write_text(
    "\n".join(lines), encoding="utf-8"
)
print("paras", len(doc.paragraphs), "lines", len(lines))
