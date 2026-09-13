from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

DOC = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题4(1).docx"
)
doc = Document(str(DOC))
lines = [f"paras={len(doc.paragraphs)} tables={len(doc.tables)}"]
for i, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    drawings = para._p.findall(".//" + qn("w:drawing"))
    extra = f"  [img={len(drawings)}]" if drawings else ""
    if t or extra:
        lines.append(f"{i:04d}{extra} {t[:220]}")
out = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q4_docx_dump.txt")
out.write_text("\n".join(lines), encoding="utf-8")
print("paras", len(doc.paragraphs), "tables", len(doc.tables), "lines", len(lines))
