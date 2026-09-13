from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

DOC = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-前言-基础部分润色稿.docx"
)
doc = Document(str(DOC))
lines = []
for i, para in enumerate(doc.paragraphs):
    n_om = para._p.xml.count("oMath")
    n_img = len(para._p.findall(".//" + qn("w:drawing")))
    t = para.text.replace("\n", " ")
    extra = []
    if n_om:
        extra.append(f"om={n_om}")
    if n_img:
        extra.append(f"img={n_img}")
    tag = (" [" + ",".join(extra) + "]") if extra else ""
    style = para.style.name if para.style is not None else ""
    lines.append(f"{i:04d} ({style}){tag} {t}")
out = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_front_dump.txt")
out.write_text("\n".join(lines), encoding="utf-8")
print("paras", len(doc.paragraphs), "tables", len(doc.tables))
