"""Rewrite Q4 8.2.3–8.2.5: brief search-then-clear, directional corrector, path opt."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订.docx"
ALT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订2.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题4(1).docx"
)
DUMP = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q4_825_after.txt")


def blank_after(anchor, style_para: Paragraph) -> Paragraph:
    new_elm = deepcopy(style_para._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    anchor._p.addnext(new_elm)
    return Paragraph(new_elm, style_para._parent)


def strip_omath(para: Paragraph) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def set_plain(para: Paragraph, text: str) -> None:
    strip_omath(para)
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def find_start(doc: Document, prefix: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def add_body(anchor, style: Paragraph, text: str, *, center: bool = False) -> Paragraph:
    p = blank_after(anchor, style)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p


def drop(para: Paragraph) -> None:
    para._p.getparent().remove(para._p)


def next_p(para: Paragraph) -> Paragraph | None:
    nxt = para._p.getnext()
    while nxt is not None and nxt.tag != qn("w:p"):
        nxt = nxt.getnext()
    if nxt is None:
        return None
    return Paragraph(nxt, para._parent)


def main() -> None:
    doc = Document(str(OUT))
    body = find_start(doc, "求解对应默认入口")
    h823 = find_start(doc, "8.2.3")
    h824 = find_start(doc, "8.2.4")
    h825 = find_start(doc, "8.2.5")
    cap5 = find_start(doc, "图 5")
    img5 = Paragraph(cap5._p.getprevious(), cap5._parent)
    note5 = find_start(doc, "覆盖四段结束后")
    steps = find_start(doc, "Step 1")
    contrast = find_start(doc, "对照分支一律不作为本文默认")

    bridge = find_start(doc, "上节给出与问题三的对照")
    old_a = find_start(doc, "对仅听过一次的频道")
    old_b = find_start(doc, "获得")
    old_c = find_start(doc, "覆盖期廉价顺路清采用")
    old_e = find_start(doc, "且本问仅当")
    old_d = next_p(old_c)
    old_batch = find_start(doc, "覆盖结束后，在估计服务点上构造无跳点开放路")
    old_svc = find_start(doc, "单步服务规则")

    set_plain(
        body,
        "求解对应默认入口 hexbatch。源位置与定向朝向均先验未知，环境以协议黑盒给出，"
        "故采用构造性求解：先由覆盖几何与动作时间确定双环检测点，再相对问题三改写定向可听条件与服务门控；"
        "随后沿用先搜后清相位，展开定向前瓣校正与路径优化，并给出编号步骤。"
        "不以整数规划或开环轨迹优化作为主求解器。",
    )
    set_plain(h823, "8.2.3  先搜后清相位")
    set_plain(h824, "8.2.4  定向前瓣校正")
    set_plain(h825, "8.2.5  路径优化与求解步骤")
    set_plain(
        note5,
        "覆盖四段结束后进入锁定开放路批量清除与扫尾。先搜后清的机理见问题三；本图只标出本问的覆盖—服务分界。",
    )

    for para in (bridge, old_a, old_b, old_c, old_d, old_e, old_batch, old_svc):
        if para is not None and para._p.getparent() is not None:
            drop(para)

    cur = add_body(
        h823,
        body,
        "覆盖—服务两阶段沿用问题三：覆盖期除近场指示与合法顺路外，不专程折返清除；"
        "听点骨架走完或已听数目达到局内源数后，再转入批量清除。"
        "“闻及即清”会把定位折返叠进外环主干，问题三已说明其时间结构差于先搜后清，此处不再展开对照。",
    )
    cur = add_body(
        cur,
        body,
        "本问只补充与定向航路有关的三点。其一，从原点前往内顶点且该顶点半径大于 940 m 时，"
        "可在射线上 900 m 停听，把近心朝外源摊进必经路程；900 m 不是第三圈任务环，也不计入完备性核验听点集。"
        "其二，听满局内源数 n 后，外十二边形停靠只保留航路上的二测频道，不再为未知信道满扫。"
        "其三，near 当场清除。图 5 标出覆盖结束后才进入锁定开放路的相位。",
    )
    cur._p.addnext(note5._p)
    cur._p.addnext(cap5._p)
    cur._p.addnext(img5._p)

    cur = add_body(
        h824,
        body,
        "定向源的示向仍按问题一构造角扇外包络 R_AOA，半宽取题面闭区间 δ=1.0°，"
        "与问题三实现余量 1.01° 区分书写。校正器只在该外包络上做朝向可行收缩，"
        "不是最小二乘或极大似然点估计，也不替换问题一的集合估计。图 4a 左下给出收缩示意。",
    )
    cur = add_body(
        cur,
        body,
        "称候选源位 G 朝向可行，若存在朝向 ψ，使全部已听测站落入 G 的前瓣闭半平面，"
        "且所有与 G 距离不超过 1000 m 的无信号点都不落入该前瓣。"
        "否则“附近无信号却本应被前瓣看见”会把错误朝向写进可行域。"
        "实现上把 ψ 离散为 720 箱（步长 0.5°）。记已听测站为 S+、无信号点为 S−，则",
    )
    cur = add_body(
        cur,
        body,
        "F(G)=1  ⇔  ∃ψ，∀P∈S+ 有 P∈H_ψ(G)，且 ∀Q∈S−，||Q−G||≤1000 m ⇒ Q∉H_ψ(G)。",
        center=True,
    )
    cur = add_body(
        cur,
        body,
        "收缩后的定位区取",
    )
    cur = add_body(cur, body, "R_dir={G∈R_AOA : F(G)=1}。", center=True)
    cur = add_body(
        cur,
        body,
        "在外包络顶点、边四分点与内部网格上取样，保留 F=1 的点，再作凸包并计算最小包围圆。"
        "出现下列情形之一则回退问题一基线：测站少于 2；外包络空、无界或顶点不足 3；"
        "可行样本少于 3；校正后最小包围圆半径大于问题一基线。"
        "近共线时即使完成收缩仍禁止清除。无信号只从该频道可行集挖洞，不把频道标成全局不存在。",
    )
    cur = add_body(
        cur,
        body,
        "可清判据仍比较最小包围圆半径与 20 m，并附加圆心朝向可行：",
    )
    cur = add_body(
        cur,
        body,
        "can_clear_20  ⇔  r_SEC(R_dir)≤20 m 且 F(c_SEC)=1。",
        center=True,
    )
    cur = add_body(
        cur,
        body,
        "仅一次示向时，第二站在问题二紧凑候选中取距当前位置最近、且前瓣兼容者；"
        "避开 35 m 内无信号点，只走一侧，不强制背面双正交。"
        "两侧均失败则沿示向取代理点 ρ≈380 m。图 4b 给出该几何。"
        "朝向箱为常数量级，外包络采样规模与定位区直径成比例；回退规则保证收缩后的 SEC 不差于问题一。",
    )

    cur = add_body(
        h825,
        body,
        "覆盖期内两项有条件机制只改变访问，不改 7+12 听点集。"
        "对仅听过一次的频道，若未来航点已落入问题二候选带（侧向 |y|≥200 m）且前瓣兼容，"
        "则在该航点顺路补第二次示向，不另插边；问题三的环边离散补测未原样拷到本问。"
        "若已有不少于两次示向且定位区可清，记当前位置为 P、下一听点为 Q、候选清除点为 C，增量路程",
    )
    cur = add_body(cur, body, "ΔL=d(P,C)+d(C,Q)−d(P,Q) ≤ 280 m", center=True)
    cur = add_body(
        cur,
        body,
        "且 C 必须是已满足 20 m 的 SEC 圆心时，才允许途清，不用示向交点冒进。"
        "该门控只收“已经能清、绕路很短”的源，避免覆盖期星形折返。",
    )
    cur = add_body(
        cur,
        body,
        "覆盖结束后，在估计服务点上构造无跳点开放路。"
        "记当前位置为 P、待清服务点集合为 C。先取局部簇 C_loc={C∈C : d(P,C)≤650 m}，"
        "在簇内按最近钉住的开放 TSP 走完，再从簇末点出发走远簇。"
        "源数不超过 16 时开放路可用精确动态规划；更大则退到最近邻加 2-opt。"
        "若局部簇尚未走完而队首已指向远点，或突然出现更近源且优势超过 80 m，则重解。"
        "这样避免欧氏最近邻把近源留到最后。",
    )
    cur = add_body(
        cur,
        body,
        "对队首频道做单步服务后即重规划。仅 1 次示向时，取最近前瓣兼容紧凑第二站（只一侧），"
        "失败则沿示向代理；已有不少于 2 次示向且可清，则清除 SEC 圆心；"
        "否则试最近示向交点，再光学网格扫描可行多边形（格距 25 m，最多 512 格，格心须前瓣兼容）。"
        "爬行仅作交会退化兜底，不是主路径。批量清后若仍有 pending，再扫尾回家清除。",
    )
    cur = add_body(
        cur,
        body,
        "编号步骤如下。Step 1　/enter，读取 jammer_count 作为停搜上界 n；忽略缺失的全向/定向个数。"
        "Step 2　原点对频道 1–20 全扫；near 当场清除。"
        "Step 3　内正七边形巡游：按需 900 m 途听、顶点扫描、顺路二测与廉价途清。"
        "Step 4　外正十二边形最近邻巡游；听满 n 后外点只做顺路二测。"
        "Step 5　覆盖结束。"
        "Step 6　近场 650 m 无跳点开放路批量清 + 扫尾。"
        "Step 7　/exit。真实时间用剩余时长减预留。图 6、图 7 为演练日志反演的稀源与密源路径。",
    )
    drop(steps)

    saved_to = None
    for dest in (OUT, ALT):
        try:
            doc.save(str(dest))
            saved_to = dest
            break
        except PermissionError:
            continue
    if saved_to is None:
        raise PermissionError("cannot save revised docx")
    copies = {"saved": str(saved_to.name)}
    for dest, key in ((WX, "wechat"), (OUT, "修订"), (ALT, "修订2")):
        if dest == saved_to:
            copies[key] = True
            continue
        try:
            copy2(saved_to, dest)
            copies[key] = True
        except OSError:
            copies[key] = False

    d2 = Document(str(saved_to))
    lines = [f"paras={len(d2.paragraphs)} copies={copies}"]
    keys = (
        "8.2",
        "求解对应",
        "8.2.3",
        "8.2.4",
        "8.2.5",
        "图 5",
        "图 6",
        "图 7",
        "覆盖—服务",
        "本问只补充",
        "定向源的示向",
        "称候选源位",
        "覆盖期内两项",
        "覆盖结束后，在估计",
        "对队首频道",
        "编号步骤",
        "对照分支",
    )
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        if t.startswith(keys) or t.startswith("图 5") or t.startswith("覆盖四段"):
            lines.append(f"{i:04d} {t[:190]}")
    DUMP.write_text("\n".join(lines), encoding="utf-8")
    print("saved", saved_to)
    print("copies", copies)
    print("\n".join(lines))


if __name__ == "__main__":
    main()
