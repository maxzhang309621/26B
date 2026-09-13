"""Fix duplicated 7.2.1 wrap paragraph after OMML leftover."""
from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.oxml.ns import qn

DOC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace\2026B-问题3-修订2.docx")
ORIG = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(1).docx"
)
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
EXTRACT = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\q3_paper_revised_extract.txt")

TEXT = (
    "综合表 2：正六边形是能够覆盖的最稀疏等角环，且在最小可行半径下覆盖巡游总时最低。"
    "实现上取 1150 m，用约 11.5 m 余量换稳健。"
    "相位不改变覆盖证书，只固定环边朝向，便于环边采样与问题二候选带判定复现。"
    "最终听点为原点加正六边形 6×1150 m（相位 10°），分布如图 2。"
    "开局仍在原点对 1–20 全扫，不计行驶时"
)


def main() -> None:
    doc = Document(str(DOC))
    para = doc.paragraphs[29]
    for tag in (f"{MATH}oMathPara", f"{MATH}oMath"):
        for el in list(para._p.findall(f".//{tag}")):
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)
    if para.runs:
        para.runs[0].text = TEXT
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(TEXT)
    doc.save(str(DOC))
    try:
        copy2(DOC, ORIG)
        copied = True
    except OSError:
        copied = False
    lines = []
    d2 = Document(str(DOC))
    for i in (19, 26, 27, 29, 33, 46, 56, 57, 69):
        lines.append(f"{i:04d} {d2.paragraphs[i].text[:220]}")
    EXTRACT.write_text("\n".join(lines), encoding="utf-8")
    print("ok copied", copied)


if __name__ == "__main__":
    main()
