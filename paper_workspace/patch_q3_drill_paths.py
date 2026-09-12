"""Replace local mock path figures with drill-log inversion maps."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

SRC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订7.docx"
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订8.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(2).docx"
)
ALT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订.docx"
FIG_DIR = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\figures")
FIG16 = FIG_DIR / "p3-20260912-130724_path.png"  # 16 sources
FIG13 = FIG_DIR / "p3-20260912-135118_path.png"  # 13 sources
FIG11 = FIG_DIR / "p3-20260912-135150_path.png"  # 11 sources
WIDTH = 13.6


def blank_after(anchor, style_para: Paragraph) -> Paragraph:
    new_elm = deepcopy(style_para._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    if hasattr(anchor, "_p"):
        anchor._p.addnext(new_elm)
    else:
        anchor.addnext(new_elm)
    return Paragraph(new_elm, style_para._parent)


def strip_omath(para: Paragraph) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def set_plain(para: Paragraph, text: str) -> None:
    strip_omath(para)
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def replace_picture(para: Paragraph, png: Path, width_cm: float) -> None:
    p_pr = para._p.find(qn("w:pPr"))
    for child in list(para._p):
        if child is not p_pr:
            para._p.remove(child)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(png), width=Cm(width_cm))


def find_start(doc: Document, prefix: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def find_caption_and_image(doc: Document, prefix: str) -> tuple[Paragraph, Paragraph]:
    prev = None
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            if prev is None:
                raise RuntimeError(f"no image before {prefix}")
            return prev, para
        prev = para
    raise RuntimeError(prefix)


def main() -> None:
    for png in (FIG16, FIG13, FIG11):
        if not png.exists():
            raise FileNotFoundError(png)

    copy2(SRC, OUT)
    doc = Document(str(OUT))

    img8, cap8 = find_caption_and_image(doc, "图 8  ")
    img9, cap9 = find_caption_and_image(doc, "图 9  ")
    p83 = find_start(doc, "表 4 按源数只列三类代表局")
    p87 = find_start(doc, "少数局的路径形态见图")

    replace_picture(img8, FIG16, WIDTH)
    replace_picture(img9, FIG11, WIDTH)

    set_plain(cap8, "图 8  演练日志反演：十六源局航迹（最多情形）")
    set_plain(cap9, "图 9  演练日志反演：十一源局航迹（少数情形）")
    set_plain(
        p83,
        "表 4 按源数只列三类代表局。少数取本批最稀的十一源，秒/源 346、总时约 3806 s，"
        "六边形骨干成本几乎不随源数下降。多数取出现次数最多的十三源（5/15 局），代表局秒/源 297，"
        "接近全样本 295 s/源。最多取十六源上界，代表局秒/源 242、总时约 3866 s，与十一源局总时接近，"
        "表明新增源主要提高骨干利用率，而非按比例拉长总时。十六源局的演练反演航迹如图 8："
        "听点环仍被完整遍历，清除折返集中于源较密扇区。十三源多数情形见图 8（续）。"
        "图 8、图 8（续）取同策略演练动作日志，源位由清除点与示向反演还原，仅作空间形态说明，"
        "不代替表 4 的官方时间口径。",
    )
    set_plain(
        p87,
        "少数局的路径形态见图 9：十一源演练反演中源周向较散、场内留白更大，覆盖骨干依旧完整，"
        "批量阶段跨空白区的边长增加，秒/源与官方十一源局同处偏高量级。"
        "稀疏布局主要恶化单位源耗时，而非削弱全清能力；不宜依据源数动态删减听点，以免破坏最不利覆盖保证。",
    )

    note = blank_after(cap8, p87)
    cap_cont = blank_after(cap8, cap8)
    img_cont = blank_after(cap8, cap8)
    img_cont.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_cont.add_run().add_picture(str(FIG13), width=Cm(WIDTH))
    cap_cont.add_run("图 8（续）  演练日志反演：十三源局航迹（多数情形）")
    note.add_run(
        "源位置由成功清除点及同局示向反演还原。三图分别对应最多、多数、少数源数，策略与表 4 官方批次相同；"
        "图中标注的虚拟时间来自作图所用演练日志，官方时间口径仍以表 4 为准。"
    )

    doc.save(str(OUT))
    copies = {}
    for dest, key in ((WX, "wechat"), (ALT, "修订")):
        try:
            copy2(OUT, dest)
            copies[key] = True
        except OSError:
            copies[key] = False

    d2 = Document(str(OUT))
    lines = [f"paras={len(d2.paragraphs)} copies={copies}"]
    prev = None
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        drawings = para._p.findall(".//" + qn("w:drawing"))
        extra = f"  [drawing={len(drawings)}]" if drawings else ""
        if any(
            k in t
            for k in (
                "图 8",
                "图 9",
                "本地仿真",
                "演练日志反演",
                "十六源",
                "十一源",
                "十三源",
                "表 4 按源数",
                "少数局的路径",
            )
        ) or (drawings and prev is not None and ("图 8" in prev or "图 9" in prev or "表 4 按源数" in prev)):
            lines.append(f"{i:04d}{extra} {t[:220]}")
        prev = t
    dump = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q3_drill_paths.txt")
    dump.write_text("\n".join(lines), encoding="utf-8")
    print("saved", OUT.name)
    print("copies", copies)
    print("\n".join(lines))


if __name__ == "__main__":
    main()
