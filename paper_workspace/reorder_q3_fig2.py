from pathlib import Path
from shutil import copy2

from docx import Document

DOC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订3.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(2).docx"
)

doc = Document(str(DOC))
hist_cap = hist_img = hist_note = fig2_note = None
for para in doc.paragraphs:
    t = para.text.strip()
    if t.startswith("图 2（续）"):
        hist_cap = para
        nxt = para._p.getnext()
        if nxt is not None:
            hist_img = nxt
            nxt2 = nxt.getnext()
            if nxt2 is not None:
                hist_note = nxt2
    if t.startswith("左图为本文采用的原点加正六边形"):
        fig2_note = para

if not all([hist_cap is not None, hist_img is not None, hist_note is not None, fig2_note is not None]):
    raise RuntimeError(f"missing {[hist_cap, hist_img, hist_note, fig2_note]}")

anchor = fig2_note._p
anchor.addnext(hist_note)
anchor.addnext(hist_cap._p)
anchor.addnext(hist_img)

doc.save(str(DOC))
try:
    copy2(DOC, WX)
    copied = True
except OSError:
    copied = False

d2 = Document(str(DOC))
lines = []
for i, para in enumerate(d2.paragraphs):
    t = para.text.strip()
    xml = para._p.xml
    mark = " [IMG]" if ("blip" in xml or "drawing" in xml) and "oMath" not in xml[:200] else ""
    if 22 <= i <= 36 or t.startswith("图 2") or t.startswith("左图") or t.startswith("各环"):
        lines.append(f"{i:04d}{mark} {t[:90]}")
Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q3_order.txt").write_text(
    "\n".join(lines), encoding="utf-8"
)
print("copied", copied)
