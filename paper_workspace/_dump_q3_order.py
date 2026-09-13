from pathlib import Path
from docx import Document

DOC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订3.docx"
doc = Document(str(DOC))
lines = []
for i, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    xml = para._p.xml
    has_img = "a:blip" in xml or "w:drawing" in xml
    mark = " [IMG]" if has_img else ""
    if 7 <= i <= 40:
        lines.append(f"{i:04d}{mark} {t[:140]}")
    if "证书" in t or "表 2" in t or "覆盖完备" in t:
        lines.append(f"HIT {i:04d} {t[:160]}")
Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q3_order2.txt").write_text(
    "\n".join(lines), encoding="utf-8"
)
print("paras", len(doc.paragraphs), "tables", len(doc.tables))
