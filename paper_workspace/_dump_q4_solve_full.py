from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

doc = Document(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace\2026B-问题4(1)-修订.docx")
lines = [f"paras={len(doc.paragraphs)}"]
for i, para in enumerate(doc.paragraphs):
    if i < 19 or i > 80:
        continue
    t = para.text.strip()
    n_img = len(para._p.findall(".//" + qn("w:drawing")))
    extra = f" [img={n_img}]" if n_img else ""
    if t or extra:
        lines.append(f"{i:04d}{extra} {t}")
Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q4_solve_full.txt").write_text(
    "\n".join(lines), encoding="utf-8"
)
print("wrote", len(lines))
