"""Rewrite Q4 8.2.5: DP principle + schematic; drop step list and inversion maps."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

SRC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订2.docx"
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订3.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题4(1).docx"
)
ALT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订.docx"
FIG = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\figures\q4_open_path_dp.png")
DUMP = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q4_825_dp.txt")


def blank_after(anchor, style_para: Paragraph) -> Paragraph:
    new_elm = deepcopy(style_para._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    anchor._p.addnext(new_elm)
    return Paragraph(new_elm, style_para._parent)


def strip_omath(para: Paragraph) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def set_plain(para: Paragraph, text: str) -> None:
    strip_omath(para)
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def find_start(doc: Document, prefix: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def add_body(anchor, style: Paragraph, text: str, *, center: bool = False) -> Paragraph:
    p = blank_after(anchor, style)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p


def drop(para: Paragraph) -> None:
    parent = para._p.getparent()
    if parent is not None:
        parent.remove(para._p)


def next_p(para: Paragraph) -> Paragraph | None:
    nxt = para._p.getnext()
    while nxt is not None and nxt.tag != qn("w:p"):
        nxt = nxt.getnext()
    if nxt is None:
        return None
    return Paragraph(nxt, para._parent)


def n_drawings(para: Paragraph) -> int:
    return len(para._p.findall(".//" + qn("w:drawing")))


def drop_until(start: Paragraph, stop_prefix: str) -> None:
    p = next_p(start)
    while p is not None:
        t = p.text.strip()
        if t.startswith(stop_prefix) or n_drawings(p):
            return
        nxt = next_p(p)
        drop(p)
        p = nxt


def replace_picture(para: Paragraph, png: Path, width_cm: float) -> None:
    p_pr = para._p.find(qn("w:pPr"))
    for child in list(para._p):
        if child is not p_pr:
            para._p.remove(child)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(png), width=Cm(width_cm))


def main() -> None:
    copy2(SRC, OUT)
    doc = Document(str(OUT))
    intro = find_start(doc, "求解对应默认入口")
    h825 = find_start(doc, "8.2.5")
    body = intro

    set_plain(
        intro,
        "求解对应默认入口 hexbatch。源位置与定向朝向均先验未知，环境以协议黑盒给出，"
        "故采用构造性求解：先由覆盖几何与动作时间确定双环检测点，再相对问题三改写定向可听条件与服务门控；"
        "随后沿用先搜后清相位，展开定向前瓣校正与无跳点开放路动态规划。"
        "不以整数规划或开环轨迹优化作为主求解器。",
    )
    set_plain(h825, "8.2.5  路径优化：无跳点开放路动态规划")

    cap6 = find_start(doc, "图 6")
    img6 = Paragraph(cap6._p.getprevious(), cap6._parent)
    note6 = find_start(doc, "蓝色为覆盖搜索段")
    cap7 = find_start(doc, "图 7")
    img7 = Paragraph(cap7._p.getprevious(), cap7._parent)
    note7 = find_start(doc, "该局对应本批次最快秒/源情形")
    cap8 = find_start(doc, "图 8")
    cap9 = find_start(doc, "图 9")
    drop_until(h825, "图 6")

    cur = add_body(
        h825,
        body,
        "覆盖期内的路径改动只作用在既有听点边上：未来航点若已落入问题二候选带且前瓣兼容，则顺路补第二次示向；"
        "已可清时，增量路程 ΔL=d(P,C)+d(C,Q)−d(P,Q)≤280 m 且 C 为 SEC 圆心才允许途清。"
        "二者都不对服务点做全局排序。覆盖结束后，待清对象变为估计服务点集合 C，起点钉在当前位置 P，且不必返回原点，"
        "这是开放哈密顿路，而不是闭回路旅行商问题。",
    )
    cur = add_body(
        cur,
        body,
        "若直接对全部服务点求欧氏开放最短路，全局最优可以先走远簇、把近旁孤立点留到最后，"
        "与“先清眼前”的时间结构相反，见图 6a 红虚线。"
        "本问因此加一层无跳点约束：先取局部簇，走完后再接远簇，",
    )
    cur = add_body(
        cur,
        body,
        "C_loc={C∈C : d(P,C)≤650 m}，  C_far=C \\ C_loc。",
        center=True,
    )
    cur = add_body(
        cur,
        body,
        "在 C_loc 内求从 P 出发的开放最短路，再从局部簇末点出发求 C_far 的开放最短路；"
        "每一簇的第一跳钉在最近城市。这样近场 650 m 内的源不会被远簇的全局捷径跳过。",
    )
    cur = add_body(
        cur,
        body,
        "簇内排序的算法依据是 Held–Karp 动态规划。把起点 P 固定，记 S 为已访问服务点子集、j∈S 为当前终点，状态"
        "f(S,j) 表示从 P 出发走遍 S 并停在 j 的最短路长。边界与转移为",
    )
    cur = add_body(
        cur,
        body,
        "f({j},j)=d(P,j)，  f(S,j)=min_{i∈S\\{j}} [f(S\\{j},i)+d(i,j)]。",
        center=True,
    )
    cur = add_body(
        cur,
        body,
        "最优开放路长为 min_j f(C,j)，回溯得次序 π*。局内源数不超过 16，状态规模 2^n·n 可精确求解；"
        "更大则退到最近邻加 2-opt。"
        "估计服务点随新示向移动，第二站又依赖当前位姿，故不能一次锁定开环回路："
        "只执行 π* 的第一座城市（对该频道做一次服务），再从新位姿重解，见图 6b。"
        "若局部簇尚未走完而队首已指向远点，或突然出现更近源且优势超过 80 m，则立即重解。",
    )
    cur._p.addnext(note6._p)
    cur._p.addnext(cap6._p)
    cur._p.addnext(img6._p)
    replace_picture(img6, FIG, 16.0)
    set_plain(cap6, "图 6  无跳点分层约束（左）与开放路 Held–Karp 动态规划（右）")
    set_plain(
        note6,
        "左图虚线圈出 650 m 局部簇。红虚线为纯欧氏开放路：远簇优先，近源留到最后；绿实线先走完局部簇再接远簇。"
        "右图：状态 f(S,j) 给出钉在 P 的精确开放路；只执行最优路的第一城，服务点随示向更新后再解。",
    )
    drop(img7)
    drop(cap7)
    drop(note7)
    set_plain(cap8, "图 7  反演：稀源局覆盖—清除分段路径")
    set_plain(cap9, "图 8  反演：密源局覆盖—清除分段路径")

    saved = OUT
    try:
        doc.save(str(OUT))
    except PermissionError:
        saved = Path(str(OUT).replace("修订3", "修订3b"))
        doc.save(str(saved))
    copies = {"saved": saved.name}
    for dest, key in ((WX, "wechat"), (ALT, "修订"), (SRC, "修订2")):
        try:
            copy2(saved, dest)
            copies[key] = True
        except OSError:
            copies[key] = False

    d2 = Document(str(saved))
    lines = [f"paras={len(d2.paragraphs)} copies={copies}"]
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        if t.startswith((
            "8.2.5", "求解对应", "覆盖期内的路径", "若直接对全部", "C_loc",
            "簇内排序", "f({j}", "最优开放路长", "图 6", "左图虚线圈出",
            "对照分支", "图 7", "图 8", "图 9", "8.4.2",
        )):
            lines.append(f"{i:04d} {t[:200]}")
    DUMP.write_text("\n".join(lines), encoding="utf-8")
    print("saved", saved)
    print("copies", copies)


if __name__ == "__main__":
    main()
