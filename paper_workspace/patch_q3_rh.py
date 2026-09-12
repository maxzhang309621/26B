"""Insert RH path-opt after search-then-clear; compress evolution section."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

SRC = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订5.docx"
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订6.docx"
WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(2).docx"
)
ALT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题3(2)-修订.docx"
FIG = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\figures\q3_rh_open_path.png")
DUMP = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q3_rh_after.txt")


def blank_after(anchor, style_para: Paragraph) -> Paragraph:
    new_elm = deepcopy(style_para._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    if hasattr(anchor, "_p"):
        anchor._p.addnext(new_elm)
    else:
        anchor.addnext(new_elm)
    return Paragraph(new_elm, style_para._parent)


def strip_omath(para: Paragraph) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def find_start(doc: Document, prefix: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def main() -> None:
    copy2(SRC, OUT)
    doc = Document(str(OUT))

    intro = find_start(doc, "鉴于源位置先验未知")
    intro.runs[0].text = (
        "鉴于源位置先验未知、环境以协议黑盒形式给出、且听点数目极少，本问采用构造性求解而非开环轨迹优化。"
        "求解顺序为：先由等圆覆盖确定正六边形检测点；再规定先搜后清的覆盖—服务相位；"
        "覆盖结束后以滚动时域开放最短路动态优化清除次序；然后分述环边补测、顺路清除与光学网格。"
        "边扫边清到先搜后清的演进与淘汰依据压缩见本节末。不以整数规划、遗传算法或强化学习作为主求解器。"
    )

    p32 = find_start(doc, "求解过程首先进入任务区")
    p32.runs[0].text = p32.runs[0].text.replace(
        "环边补测、顺路清除与光学网格是骨架上的三项有条件机制，分见后文。",
        "清除路径的滚动优化见 7.2.3；环边补测、顺路清除与光学网格分见其后。",
    )

    h222 = find_start(doc, "7.2.2")
    note3 = find_start(doc, "该对照说明")
    batch = find_start(doc, "当听点遍历完毕或已听数目达到局内源数后")

    h223 = blank_after(note3, h222)
    h223.add_run("7.2.3  清除路径的滚动开放路优化")

    batch.runs[0].text = (
        "覆盖结束后进入清除路径的动态优化。各信道的估计服务点随新示向移动，第二站又依赖当前位姿，"
        "一次锁定的开环旅行商回路会在下一步观测后失效。因此采用滚动时域开放最短路："
        "以当前位置 P 为起点，对尚未清除的估计服务点 {c_i} 求开放哈密顿路，只执行路的第一座城市对应的一个协议动作，"
        "再从新位姿重解。服务点取法：两站交会且最小包围圆半径不超过 20 m 时取圆心，否则取问题二推荐点；"
        "示向交点只试距离当前位置最近的一个，避免远交点造成星形折返。"
    )

    f1 = blank_after(batch, batch)
    f1.add_run("开放路目标为")
    f2 = blank_after(f1, batch)
    f2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f2.add_run("π* = arg min_π [ d(P, c_{π(1)}) + Σ_{i=1}^{m−1} d(c_{π(i)}, c_{π(i+1)}) ]。")
    f3 = blank_after(f2, batch)
    f3.add_run(
        "待清点数不超过 16 时，以当前位置为附加起点用 Held–Karp 动态规划求精确开放路，状态转移为"
    )
    f4 = blank_after(f3, batch)
    f4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f4.add_run("f(S, j) = min_{i∈S\\{j}} [ f(S\\{j}, i) + d(c_i, c_j) ]，")
    f5 = blank_after(f4, batch)
    f5.add_run(
        "其中 f({P}, P)=0；点数更多时改用最近邻加 2-opt。已有不少于两次示向的源整段直奔服务点，不再按 280 m 步进途听；"
        "仅一次示向时步长取 min{280 m, max(120 m, 0.5 d(P,c))}。"
        "换路设置 50 m 滞回：若保持当前目标的开放路长不超过新最优路长加 50 m，则不改序，以免估计点微动就振荡。"
        "前缀动作必须落在检测或清除指令上，协议不允许空驶。有界沿示向爬行仅作交会退化时的兜底，正常路径爬行次数应为 0。"
        "本地 40 种子对照：扫完后硬两波（先只排不少于两次示向）约 354 s/源，相对滚动重解加环边补测后的 301.7 s/源约高 52；"
        "关边补时约 312.6 s/源。说明真正下降来自环上补全二测后再滚动重解，而不是一次锁定访问序。图 3（续）示意锁定路与逐步重解的差别。"
    )

    cap_style = find_start(doc, "图 3  ")
    img = blank_after(f5, cap_style)
    img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img.add_run().add_picture(str(FIG), width=Cm(16.0))
    cap = blank_after(img, cap_style)
    cap.add_run("图 3（续）  锁定开放路与滚动时域重规划")
    note = blank_after(cap, batch)
    note.add_run(
        "左图一次排定 P→c2→c4→c1→c3；右图执行第一城后服务点关系改变，从新位姿重解剩余城市。"
        "每步只执行最优开放路的第一座城市。"
    )

    renames = (
        ("7.2.3  环边补测", "7.2.4  环边补测：把第二次示向摊进已有环边"),
        ("7.2.4  顺路清除", "7.2.5  顺路清除：用增量路程限制覆盖期内的有条件服务"),
        ("7.2.5  光学网格", "7.2.6  光学网格：仅在小定位区上作有界试清"),
        ("7.2.6  从边扫边清", "7.2.7  从边扫边清到先搜后清的演进与取舍"),
    )
    for old, new in renames:
        para = find_start(doc, old)
        if para.runs:
            para.runs[0].text = new
            for run in para.runs[1:]:
                run.text = ""

    mech = find_start(doc, "上述结构的机理在于")
    for run in mech.runs:
        if run.text:
            run.text = run.text.replace("7.2.6", "7.2.7")

    p58 = find_start(doc, "针对干扰源位置未知且移动代价")
    p58.runs[-1].text = (
        "，主干规划与回接亦不可忽视。因而主要矛盾是清除把机体带离覆盖主干后的绕行。"
        "时间账本见图 6。"
    )

    p62 = find_start(doc, "三类代表性探索")
    strip_omath(p62)
    p62.runs[0].text = (
        "首段入口自适应、八点收半径与服务驱动外扩等均可在均值上变快，但配对尾部回退，故未进入主线；"
        "六点环在在线清除逻辑下均值与尾部均劣于八点基准，说明短主干并不能在边扫边清中自动兑现；"
        "跨未来边插入均值变慢。环边补测与有限顺路清除被迁移到先搜后清骨架。"
        "代表性候选摘要见表 3，均值—尾部平面见图 7。本地 Mock 与官方演练不得混排。"
    )
    for run in p62.runs[1:]:
        run.text = ""

    p68 = find_start(doc, "在线小批次服务")
    p68.runs[0].text = (
        "小批次在线服务与交会角审计因触发不足，不展开。边扫边清保留为对照："
        "它说明均值改善不等于尾部稳定，并筛选出可迁移的环边机制。"
    )

    doc.save(str(OUT))
    copies = {}
    for dest, key in ((WX, "wechat"), (ALT, "修订")):
        try:
            copy2(OUT, dest)
            copies[key] = True
        except OSError:
            copies[key] = False

    d2 = Document(str(OUT))
    lines = [f"paras={len(d2.paragraphs)} copies={copies}"]
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        if t.startswith(("7.2", "图 3", "开放路", "π*", "f(S", "首段入口", "针对干扰", "小批次", "上述结构", "当听点遍历")):
            lines.append(f"{i:04d} {t[:190]}")
    DUMP.write_text("\n".join(lines), encoding="utf-8")
    print("saved", OUT)
    print("copies", copies)


if __name__ == "__main__":
    main()
