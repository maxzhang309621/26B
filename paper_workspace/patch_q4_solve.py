"""Expand Q4 8.2: dual-ring proof first, then Q3-to-Q4 directional changes."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

WX = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题4(1).docx"
)
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace") / "2026B-问题4(1)-修订.docx"
FIG = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\figures")
DUMP = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\_q4_solve_after.txt")
FIG_GEO = FIG / "q4_double_ring_geometry.png"
FIG_TIME = FIG / "q4_ngon_time_ranked.png"
FIG_ALGO = FIG / "q4_q3_vs_q4_algo.png"


def blank_after(anchor, style_para: Paragraph) -> Paragraph:
    new_elm = deepcopy(style_para._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    if hasattr(anchor, "_p"):
        anchor._p.addnext(new_elm)
    else:
        anchor.addnext(new_elm)
    return Paragraph(new_elm, style_para._parent)


def strip_omath(para: Paragraph) -> None:
    for el in list(para._p):
        if "oMath" in el.tag:
            para._p.remove(el)


def set_plain(para: Paragraph, text: str) -> None:
    strip_omath(para)
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def replace_picture(para: Paragraph, png: Path, width_cm: float) -> None:
    p_pr = para._p.find(qn("w:pPr"))
    for child in list(para._p):
        if child is not p_pr:
            para._p.remove(child)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(png), width=Cm(width_cm))


def find_start(doc: Document, prefix: str) -> Paragraph:
    for para in doc.paragraphs:
        if para.text.strip().startswith(prefix):
            return para
    raise RuntimeError(prefix)


def add_body(anchor, style: Paragraph, text: str, *, center: bool = False) -> Paragraph:
    p = blank_after(anchor, style)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p


def add_figure(anchor, body: Paragraph, cap_style: Paragraph, png: Path,
               width_cm: float, caption: str, note: str) -> Paragraph:
    img = blank_after(anchor, body)
    cap = blank_after(img, cap_style)
    note_p = blank_after(cap, body)
    replace_picture(img, png, width_cm)
    cap.add_run(caption)
    note_p.add_run(note)
    return note_p


def main() -> None:
    copy2(WX, OUT)
    doc = Document(str(OUT))

    intro = find_start(doc, "求解对应默认入口")
    h821 = find_start(doc, "8.2.1")
    h822 = find_start(doc, "8.2.2")
    h823 = find_start(doc, "8.2.3")
    h824 = find_start(doc, "8.2.4")
    p28 = find_start(doc, "设内顶点为")
    cap3 = find_start(doc, "图 3")
    cap4 = find_start(doc, "图 4")
    cap5 = find_start(doc, "图 5")
    cap6 = find_start(doc, "图 6")
    cap7 = find_start(doc, "图 7")
    img3 = Paragraph(cap3._p.getprevious(), cap3._parent)
    note3 = find_start(doc, "覆盖四段结束后")
    body = intro
    cap_style = cap3

    set_plain(h822, "8.2.3  覆盖期二测、校正与顺路清")
    set_plain(h823, "8.2.4  批量清除与单步服务")
    set_plain(h824, "8.2.5  先搜后清、算法步骤与路径形态")
    set_plain(h821, "8.2.1  双环检测点：覆盖几何与动作时间最优性")

    set_plain(cap7, "图 9  反演：密源局覆盖—清除分段路径")
    set_plain(cap6, "图 8  反演：稀源局覆盖—清除分段路径")
    set_plain(cap5, "图 7  问题四密源代表局路径（16 源）")
    set_plain(cap4, "图 6  问题四演练路径反演：稀源局与密源局对照")
    set_plain(cap3, "图 5  问题四先搜后清相位")

    set_plain(
        intro,
        "求解对应默认入口 hexbatch。源位置与定向朝向均先验未知，环境以协议黑盒给出，"
        "故采用构造性求解：先由覆盖几何与动作时间确定双环检测点，再相对问题三改写定向可听条件与服务门控，"
        "最后给出先搜后清相位与编号步骤。不以整数规划或开环轨迹优化作为主求解器。",
    )

    intro._p.addnext(h821._p)
    h824._p.addnext(note3._p)
    h824._p.addnext(cap3._p)
    h824._p.addnext(img3._p)

    cur = add_body(
        h821,
        body,
        "全向覆盖仍落在“原点加正 n 边形环”这一结构上。工作圆半径 R=1800 m，最坏有效接收半径 r_eff=1000 m。"
        "环带上距原点与环点都最远的位置出现在相邻顶点的角平分线与工作圆的交点。"
        "该点可被半径 r_eff 的等圆盖住，当且仅当外端条件成立：",
    )
    cur = add_body(cur, body, "R sin(π/n) ≤ r_eff。", center=True)
    cur = add_body(
        cur,
        body,
        "取 n=5 时，1800 sin(π/5)≈1058 m，已越出 1000 m，五点环不可行；取 n=6 时，1800 sin(π/6)=900 m ≤ 1000 m。"
        "因此全向层至少需要正六边形。图 3a 标出七边形相邻顶点的角平分线最远点，用以核对内环是否压在覆盖区间的左端。",
    )
    cur = add_body(
        cur,
        body,
        "在 n≥6 时，原点加正 n 边形的可行外接半径构成闭区间。由角平分线两端的覆盖条件，左端为",
    )
    cur = add_body(
        cur,
        body,
        "ρ_min(n)=R cos(π/n) − [r_eff^2 − R^2 sin^2(π/n)]^(1/2)，",
        center=True,
    )
    cur = add_body(
        cur,
        body,
        "右端不超过 2 r_eff cos(π/n)。覆盖巡游路程随 ρ 递增，故在可行区间内应取左端。"
        "n=7 时 ρ_min≈997.2 m，低于正六边形门槛 1123 m，周向弦也更短。"
        "本文内环即取该左端，使原点与七顶点在最坏 1000 m 下完成全向完备性核验，同时尽量压缩内环空驶。",
    )
    cur = add_body(
        cur,
        body,
        "定向源把可听从等圆改成圆盘与前瓣半平面之交。朝向圆外的源，其前瓣落在任务圆外，内环顶点落入后瓣即静默，见图 3b。"
        "因此全向可行并不蕴含定向可行：必须在圆外布置听点。"
        "外环取正十二边形、半径 1865 m，并与内环同径向对齐，使由内顶点切入外环时沿同一射线，避免交错相位多付一段弦长。"
        "粗网格核验下，7+12 同径向对全向与前瓣均通过；6+12 无途听环全向通过但前瓣失败；7+11×1900 同样前瓣失败。"
        "故外环不能减到 11 点，内环也不能退回无外环的问题三六点方案。",
    )
    cur = add_body(
        cur,
        body,
        "访问顺序上，内环按最近邻；外环从当前位置反复选取最近未访顶点。"
        "同环全向冗余可跳过，但不得用内环听点压制外环认证顶点。"
        "900 m 途听只出现在原点出发、内顶点半径大于 940 m 的射线上，专用于近心朝外源，不计入完备性核验听点集。",
    )
    cur = add_figure(
        cur,
        body,
        cap_style,
        FIG_GEO,
        16.0,
        "图 3  双环检测点的覆盖几何：内七边形角平分线（左）与朝外源前瓣（右）",
        "左图虚线为任务圆；红叉为相邻内顶点角平分线与工作圆的交点，浅盘为半径 1000 m 可听范围。"
        "右图阴影为朝外定向源的前瓣可听盘：内顶点落入后瓣，外环顶点进入前瓣，故圆外听点是定向覆盖的必要采样，而不是对问题三的简单加圈。",
    )
    cur = add_body(
        cur,
        body,
        "覆盖巡游的虚拟时间须把几何路程与协议动作一并计入。"
        "原点对 1–20 信道全扫耗时 119 s；其后每个停靠在最坏情形下仍须扫描全部未知信道，单点驻留 120 s。"
        "以 5 m/s 折合行驶时间，停靠数为 N 的先搜后清开放巡游满足",
    )
    cur = add_body(cur, body, "T = L/5 + 119 + (N−1)×120，", center=True)
    cur = add_body(
        cur,
        body,
        "其中 L 为原点经内环再到外环、不返回原点的路程。"
        "完备性核验的听点集不含 900 m 环时，7+12 方案 N=20。",
    )
    cur = add_body(
        cur,
        body,
        "在该口径下：7 内环取 ρ_min、12 外环取 1865 m 且同径向时，行驶 3655 s、驻留 2399 s，合计 6054 s，"
        "是前瓣核验通过方案中的最短。交错相位同点集行驶增至 3707 s，合计 6106 s。"
        "内环增至正八、正九或外环增至十四边形，停靠驻留立刻把总时抬到 6180–6328 s。"
        "若把 900 m 计入听点集以使正六内环通过前瓣，停靠增至 25、总时 7888 s；"
        "旧方案 8×1200+12×2100+途听环为 8773 s。"
        "7+11×1900 合计 5958 s 更短，但前瓣核验未通过，不能作为可行最短。"
        "图 3（续）按总时排序给出上述对照。",
    )
    cur = add_figure(
        cur,
        body,
        cap_style,
        FIG_TIME,
        14.6,
        "图 3（续）  双环覆盖巡游虚拟时间（最坏满扫 20 信道）",
        "柱长为行驶与驻留之和。斜线柱表示粗网格前瓣核验未通过。虚线标出本文 7+12 同径向方案。"
        "该时间为覆盖下界，不含定位清除折返；官方演练成绩见 8.3 节。",
    )

    h_new = add_body(p28, h821, "8.2.2  由问题三到问题四：定向可听与服务门控")
    cur = add_body(
        h_new,
        body,
        "问题四与问题三共用问题一的角扇交定位核与问题二的第二站接口，差别发生在“听得到”和“何时清”。"
        "问题三可听当且仅当检测点落入以源为心、半径 r_eff 的闭圆盘；问题四还要求该点落在未知朝向 ψ 的前瓣闭半平面内：",
    )
    cur = add_body(
        cur,
        body,
        "P ∈ D(G, r_eff) ∩ H_ψ，  H_ψ={P:(P−G)·u(ψ)≥0}。",
        center=True,
    )
    cur = add_body(
        cur,
        body,
        "因此覆盖层由单环等圆变为双环前瓣采样；服务层仍引用问题一、二的公式，不另写沿示向爬行作为常规定位。"
        "示向误差实现取 δ=1.0°，与问题三实现余量 1.01° 区分书写。图 4a 对照两种可听集合。",
    )
    cur = add_body(
        cur,
        body,
        "听点上，问题三锁定原点加正六边形 6×1150 m（相位 10°）；本问锁定原点加内正七边形 7×ρ_min 与外正十二边形 12×1865 m。"
        "问题三无需圆外听点，本问外环不可省。途听 900 m 是本问新增的出发顺路检测，不是问题三的第三环。",
    )
    cur = add_body(
        cur,
        body,
        "覆盖期二测亦随之改写。问题三在环边 A→B 上取 t∈{0.25,0.40,0.50,0.60,0.75} 插入合法第二站。"
        "本问外环存在跳点与前瓣门控，环边补测未原样拷贝；仅当未来航点已落入问题二候选带（侧向 |y|≥200 m）且前瓣兼容时，"
        "才在该航点顺路补听，不另插边。",
    )
    cur = add_body(
        cur,
        body,
        "定位收缩针对定向朝向。不少于两次示向后，先按问题一构造角扇外包络 R_AOA，"
        "再在 720 个朝向箱上保留“存在解释朝向、且附近无信号不与前瓣矛盾”的点，得到 R_dir ⊆ R_AOA。"
        "若可行样本少于 3、仅有单站，或校正后最小包围圆半径大于问题一基线，则回退 R_AOA。"
        "可清判据仍是 r_SEC≤20 m，并要求圆心朝向可行。校正器是收缩不是替换，见图 4a 左下。",
    )
    cur = add_body(
        cur,
        body,
        "顺路清除的增量路程门控与问题三相同，",
    )
    cur = add_body(cur, body, "ΔL=d(P,C)+d(C,Q)−d(P,Q) ≤ 280 m，", center=True)
    cur = add_body(
        cur,
        body,
        "但问题三允许在已有不少于两次示向后前往估计服务点；本问只当 C 为已满足 20 m 的 SEC 圆心时才途清，不用示向交点冒进。"
        "该机制不改变 7+12 听点完备性集合。图 4a 右下给出 ΔL 的几何含义。",
    )
    cur = add_body(
        cur,
        body,
        "第二站在全向情形下可在示向两侧给出紧凑候选；定向源背面正交常落入后瓣，本问只走前瓣兼容的一侧，"
        "两侧均失败则沿示向取代理点 ρ≈380 m，见图 4b。"
        "批量清除由问题三的滚动开放路改为近场 650 m 无跳点锁定开放路，避免欧氏最近邻把近源留到最后。",
    )
    add_figure(
        cur,
        body,
        cap_style,
        FIG_ALGO,
        16.0,
        "图 4  由问题三到问题四的算法改动：可听集合与顺路门控（左），第二站只走前瓣一侧（右）",
        "左上：等圆变为盘与前瓣半平面之交。左下：校正器在角扇外包络内保留朝向可行点；"
        "顺路清只前往已可 20 m 的 SEC 圆心。右：第一站示向后，背面正交不采用，无合法紧凑点时沿前瓣代理。",
    )

    doc.save(str(OUT))
    copies = {}
    try:
        copy2(OUT, WX)
        copies["wechat"] = True
    except OSError:
        copies["wechat"] = False

    d2 = Document(str(OUT))
    lines = [f"paras={len(d2.paragraphs)} copies={copies}"]
    keys = (
        "8.2",
        "图 3",
        "图 4",
        "图 5",
        "图 6",
        "图 7",
        "图 8",
        "图 9",
        "求解对应",
        "全向覆盖仍落",
        "覆盖巡游的虚拟",
        "问题四与问题三共用",
        "听点上，问题三",
        "第二站在全向",
    )
    for i, para in enumerate(d2.paragraphs):
        t = para.text.strip()
        if t.startswith(keys):
            lines.append(f"{i:04d} {t[:200]}")
    DUMP.write_text("\n".join(lines), encoding="utf-8")
    print("saved", OUT)
    print("copies", copies)
    print("\n".join(lines[:40]))


if __name__ == "__main__":
    main()
