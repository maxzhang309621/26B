"""Minimal Q3 paper revision: keep OMML/figures, fix skill gaps."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(1).docx"
)
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace\2026B-问题3-修订.docx")
EXTRACT = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\q3_paper_revised_extract.txt")

STEPS = (
    "求解步骤如下。第一步：进入任务，局内源数取进入接口返回值，不以 16 为缺省。"
    "第二步：在原点对 1–20 信道全扫一次。"
    "第三步：沿正六边形前进；每条边上对只听过一次、且后续顶点补不上合法二测的信道插入环边二测；"
    "顶点扫描尚未排除的信道，并在增量路程不超过 280 m 时顺路清除。"
    "第四步：覆盖航路走尽，或已听加待清达到局内源数，则结束覆盖。"
    "第五步：对剩余待清源做滚动开放最短路批量清除；待清源数不超过 12 时用 Held–Karp 精确求解，"
    "否则用最近邻加 2-opt。已有不少于两次示向者整段直奔服务点，交会只试最近一个交点。"
    "第六步：退出任务。真实时间按剩余时长减去预留控制。"
)

RH = (
    "覆盖结束后，对尚未清除的源以其估计服务点为城市，做滚动时域开放最短路："
    "待清源数不超过 12 时用 Held–Karp 精确求解，否则用最近邻加 2-opt。"
    "已有不少于两次示向的源，整段直奔最小包围圆圆心或最近示向交点，不再按 280 m 步进途听；"
    "只试距离当前位置最近的一个交点，远交点会造成星形折返。"
    "滚动重解的滞回取 50 m，避免估计点微动就改路。协议不允许空驶，相邻服务点之间仍按合法动作衔接。"
    "有界沿示向爬行仅作数值退化兜底，每信道最多 40 步、步长 12 m，正常路径爬行次数应为 0，不写入主路径。"
)


def rewrite_plain(para, text: str) -> None:
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def insert_after(paragraph, text: str) -> Paragraph:
    new_elm = deepcopy(paragraph._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    paragraph._p.addnext(new_elm)
    new_para = Paragraph(new_elm, paragraph._parent)
    run = new_para.add_run(text)
    if paragraph.runs:
        r_pr = paragraph.runs[0]._r.find(qn("w:rPr"))
        if r_pr is not None:
            run._r.insert(0, deepcopy(r_pr))
    return new_para


def delete_paragraph(paragraph) -> None:
    elm = paragraph._p
    parent = elm.getparent()
    if parent is not None:
        parent.remove(elm)


def replace_cell_text(cell, old: str, new: str) -> None:
    if old not in cell.text:
        return
    for para in cell.paragraphs:
        for run in para.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)


def extract_text(doc: Document, path: Path) -> None:
    lines = [f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}"]
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            lines.append(f"{i:04d} {text}")
    lines.append("")
    lines.append("===== TABLES =====")
    for ti, table in enumerate(doc.tables):
        lines.append(f"--- table {ti} {len(table.rows)}x{len(table.columns)} ---")
        for ri, row in enumerate(table.rows):
            cells = " | ".join(c.text.replace("\n", " ") for c in row.cells)
            lines.append(f"  r{ri}: {cells}")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    copy2(SRC, OUT)
    doc = Document(str(OUT))
    p = doc.paragraphs

    # 7.1 去掉代码腔、停搜不用默认 16
    p[9].runs[15].text = "近距返回"
    p[9].runs[20].text = "无信号回报 "
    p[10].runs[3].text = "近距返回"
    p[10].runs[5].text = "待清源"
    p[10].runs[8].text = "（不以默认"
    p[10].runs[10].text = "作为分母），或覆盖航路走尽。三阶段示意如图"
    rewrite_plain(
        p[13],
        "覆盖期以听全为主，边补二测与顺路清只服务“不破坏骨干”的机会；"
        "批量清除期调用问题一、二；收尾保留扇形补清与收紧后的光学网格，"
        "正常路径沿示向爬行次数应为 0。",
    )

    # 7.2 四条主线 + 步骤
    rewrite_plain(
        p[17],
        "环境为黑盒接口、源位置未知、听点个数极少，故求解取构造性策略而非开环轨迹优化。"
        "本节按四条主线展开：先用覆盖证书固定正六边形参数；再规定光学网格与清除失败的启用条件；"
        "然后给出环边补测与顺路清除规则；最后在覆盖结束后用滚动开放最短路批量清除。",
    )
    p[34].runs[2].text = "清除失败（miss）"
    p[35].runs[9].text = "清除指令"
    p[45].runs[1].text = "先搜后清、环边补测、顺路清除与滚动开放路"
    p[51].runs[9].text = "待清源"
    p[55].runs[3].text = "清除折返"
    p[55].runs[25].text = "开放最短路"
    rewrite_plain(
        p[59],
        "因此默认相位取先搜后清：覆盖期除近距返回与廉价顺路清外不专程离开覆盖骨干去清除；"
        "专程清除留到听点走完后开放路排序。环边补测与顺路清是配套机制，不是听完即清。",
    )
    p[60].runs[13].text = "），样点落入问题二候选带则插入二测；后续顶点已能提供合法二测则不插。"
    p[63].runs[4].text = "发出清除指令"
    rewrite_plain(p[67], RH)

    # 7.3 标题与机制计数
    p[68].runs[1].text = "结果分析"
    p[71].runs[5].text = "边补二测批次（20260912）"
    p[71].runs[16].text = "进入接口给出的局内源数"
    p[71].runs[37].text = (
        "；表中行驶分解取自本地日志与官方结果对齐后的骨干扫描、定位行驶与清除折返。"
        "该批次 15 局环边复测合计约 150 次、顺路清合计约 50 次。"
    )
    p[91].runs[18].text = (
        "源抬升”。15 局环边复测合计约 150 次、顺路清合计约 50 次。"
        "光学网格因启用门槛高，对均时贡献有限，其价值在于抑制大角扇空跑。"
    )

    # 7.4 边补合法性、去掉代码腔
    p[104].runs[7].text = "问题一定位质量判定"
    p[108].runs[1].text = "对照"
    p[108].runs[5].text = "待清源"
    p[108].runs[6].text = (
        "清空；边补插入点均落在问题二候选带内，后续顶点已能二测则不插；非法动作为零；官方"
    )
    p[108].runs[9].text = "局内源数"
    p[118].runs[23].text = "本地对照"
    p[118].runs[46].text = (
        "，提交口径锁定为六边形先搜后清加边补二测、顺路清除与收紧光学网格。"
        "边补样点必须落入问题二候选带；若后续顶点已能提供合法二测，则不在边上插入。"
    )

    # 7.5 并入 7.4 收束，避免第五个并列大节
    p[120].runs[0].text = "综上，" + p[120].runs[0].text
    p[120].runs[7].text = " 清除失败 "
    p[120].runs[11].text = (
        "顺路清除，覆盖结束后再以滚动开放最短路批量收口；定位核始终调用问题一、二。"
        "选定先搜后清，是因为听完即清会撕碎覆盖骨干，而扫完后再拧顺序空间很小，折返须在环上砍掉。官方演练"
    )

    insert_after(p[17], STEPS)

    for para in list(doc.paragraphs):
        if para.text.replace(" ", "").startswith("7.5") and "总结" in para.text:
            delete_paragraph(para)
            break

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_cell_text(cell, "cleared=jammer_count", "清除数=局内源数")
                replace_cell_text(cell, "jammer_count", "局内源数")

    doc.save(str(OUT))
    extract_text(Document(str(OUT)), EXTRACT)
    print(f"saved {OUT}")
    print(f"extract {EXTRACT}")

    try:
        copy2(OUT, SRC)
        print(f"also wrote {SRC}")
    except OSError as exc:
        print(f"original locked/unwritable: {exc}")


if __name__ == "__main__":
    main()
