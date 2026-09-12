"""Restructure Q3 7.2: hexagon proof, search-then-clear evidence, RH path opt."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copy2
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image

REV = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace\2026B-问题3-修订.docx")
OUT = Path(r"D:\maxzhang\python files\Math-modeling\26B\paper_workspace\2026B-问题3-修订2.docx")
ORIG = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题3(1).docx"
)
FIG = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\figures\q3_hexagon_waypoints.png")
EXTRACT = Path(r"D:\maxzhang\python files\Math-modeling\26B\output\q3_paper_revised_extract.txt")
MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"

HEX_PROOF = (
    "因此在“原点加正多边形环”这一结构下，正六边形是能够覆盖的最少停靠方案，也是最小可行半径下覆盖巡游总时最短的方案。"
    "表 2 把各环取到刚能覆盖的最小半径：五点环因 "
    "1800 sin(π/5) > 1000 m 越界；六、七、八边形均可覆盖，巡游总时分别约 2187 s、2197 s、2272 s，"
    "多一个顶点虽能缩短环向行驶，但每个环点最坏要多付一轮未知信道扫描驻留，总时上升。"
    "本文环半径取 1150 m，略高于六边形门槛 1123 m，最坏最近距离约 988.5 m，余量约 11.5 m，"
    "覆盖巡游总时约 2219 s，仍低于七、八边形的最小半径方案。相位 10° 不改变覆盖证书，只固定环边相对坐标轴的朝向，便于复现。"
)

RH_TEXT = (
    "覆盖结束后进入清除路径的动态优化。未知源的估计服务点随新示向移动，第二站又依赖当前位姿，故不能一次锁定开环旅行商回路，"
    "而采用滚动时域开放最短路：每走完一步，从新位置对尚未清除的源重解开放路，只执行路的第一座城市。"
    "待清源数不超过 12 时用 Held–Karp 动态规划求精确开放路，否则用最近邻加 2-opt。"
    "已有不少于两次示向的源整段直奔最小包围圆圆心或最近示向交点，不再按 280 m 步进途听；交会只试最近一个交点。"
    "换路滞回取 50 m，避免估计点微动就改序。协议不允许空驶。"
    "表 3 给出同分布 40 种子对照：硬两波（先只排不少于两次示向）秒/源约 354，相对本文 301.7 约高 52；"
    "软偏置约 319–326，未就绪改去补测点约 346，扫完后再宽松顺路清约 314，均净亏。"
    "真正下降来自环上补全二测后再滚动重解——边补打开后秒/源由 312.6 降至 301.7，环上定位行驶由约 475 s 降至 222 s。"
    "有界沿示向爬行仅作退化兜底，正常路径爬行次数应为 0。"
)


def rewrite_plain(para, text: str) -> None:
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run.text = ""
    else:
        para.add_run(text)


def strip_omath(para) -> None:
    for tag in (f"{MATH}oMathPara", f"{MATH}oMath"):
        for el in list(para._p.findall(f".//{tag}")):
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)


def insert_after(after_para, text: str, style_para=None) -> Paragraph:
    src = style_para or after_para
    new_elm = deepcopy(src._p)
    p_pr = new_elm.find(qn("w:pPr"))
    for child in list(new_elm):
        if child is not p_pr:
            new_elm.remove(child)
    after_para._p.addnext(new_elm)
    new_para = Paragraph(new_elm, after_para._parent)
    run = new_para.add_run(text)
    if src.runs:
        r_pr = src.runs[0]._r.find(qn("w:rPr"))
        if r_pr is not None:
            run._r.insert(0, deepcopy(r_pr))
    return new_para


def set_cell(cell, text: str) -> None:
    paras = cell.paragraphs
    if not paras:
        return
    if paras[0].runs:
        paras[0].runs[0].text = text
        for run in paras[0].runs[1:]:
            run.text = ""
    else:
        paras[0].add_run(text)
    for para in paras[1:]:
        for run in para.runs:
            run.text = ""


def extract_text(doc: Document, path: Path) -> None:
    lines = [f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}"]
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            lines.append(f"{i:04d} {text}")
    lines.append("")
    lines.append("===== TABLES =====")
    for ti, table in enumerate(doc.tables):
        lines.append(f"--- table {ti} {len(table.rows)}x{len(table.columns)} ---")
        for ri, row in enumerate(table.rows):
            cells = " | ".join(c.text.replace("\n", " ") for c in row.cells)
            lines.append(f"  r{ri}: {cells}")
    path.write_text("\n".join(lines), encoding="utf-8")


def replace_fig2(docx_path: Path, png_path: Path) -> None:
    im = Image.open(png_path)
    w_px, h_px = im.size
    old_cx = 5669280
    new_cy = int(round(old_cx * h_px / w_px))
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/media/image2.png":
                data = png_path.read_bytes()
            elif item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                xml = xml.replace('cx="5669280" cy="2741899"', f'cx="{old_cx}" cy="{new_cy}"')
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(docx_path)


def main() -> None:
    doc = Document(str(REV))
    p = doc.paragraphs

    rewrite_plain(
        p[17],
        "环境为黑盒接口、源位置未知、听点个数极少，故求解取构造性策略而非开环轨迹优化。"
        "本节按三条主线展开：用覆盖证书证明正六边形检测点最优，并只给出其分布图；"
        "用增量路程与演练对照证明先搜后清优于边扫边清；"
        "覆盖结束后用滚动开放最短路动态优化清除次序，并给出相对硬两波、软偏置等对照的提升。"
        "光学网格仅作为小定位区上的有界试清，规则见 7.2.2。",
    )
    p[19].runs[1].text = "正六边形检测点及其几何最优性"

    p[25].runs[19].text = (
        "直接否决。可行情形须 n≥6；再在最小可行半径下比较覆盖巡游总时，正六边形最短，见表 2。"
    )

    strip_omath(p[26])
    rewrite_plain(p[26], HEX_PROOF)

    p[27].runs[0].text = "表 2  原点加正 n 边形覆盖巡游对照（最坏接收半径 1000 m）"
    for run in p[27].runs[1:]:
        run.text = ""

    p[29].runs[0].text = "综合表"
    p[29].runs[2].text = (
        "：正六边形是能够覆盖的最稀疏等角环，且在最小可行半径下覆盖巡游总时最低。"
        "实现上取略高于门槛的"
    )
    p[29].runs[8].text = "。解析最坏距离对整体旋转不变，相位"
    p[29].runs[18].text = "），检测点分布如图"

    rewrite_plain(
        p[31],
        "其作用是把未知信道集压到“原点听不到的子集”；其后环点只扫该子集。"
        "第二次示向尽量摊进六边形环边，进一步放大少停靠带来的折返收益。",
    )
    rewrite_plain(p[33], "图 2  正六边形检测点分布：原点全扫与环上 6×1150 m（相位 10°）")
    rewrite_plain(
        p[34],
        "虚线为工作圆 1800 m；浅盘为最坏接收半径 1000 m 的可听范围，画在 1 号环点上示意。"
        "该图只展示本文采用的六边形检测点，最优性由覆盖证书与表 2 的巡游总时给出，不再并列八边形分布。",
    )

    table = doc.tables[1]
    headers = ["方案", "覆盖", "环半径/m", "巡游总时/s", "驻留/s", "环向行驶/s"]
    rows = [
        ["五点环", "否", "—", "—", "—", "—"],
        ["正六边形（最小半径）", "是", "1123", "2187", "839", "1348"],
        ["正七边形（最小半径）", "是", "997", "2197", "959", "1238"],
        ["正八边形（最小半径）", "是", "938", "2272", "1079", "1193"],
    ]
    for j, h in enumerate(headers):
        set_cell(table.rows[0].cells[j], h)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            set_cell(table.rows[i].cells[j], val)
    extra = table.add_row()
    extra_vals = ["正六边形（本文）", "是", "1150", "2219", "839", "1380"]
    for j, val in enumerate(extra_vals):
        set_cell(extra.cells[j], val)

    p[46].runs[1].text = "先搜后清与边扫边清"
    rewrite_plain(
        p[47],
        "相位上要在“先走完检测听点、再批量清除”（先搜后清）与“每听一源立刻专程清除”（边扫边清）之间取舍。"
        "下面用增量路程公式、路径示意与演练时间账本给出证据。",
    )
    p[56].runs[24].text = (
        "清除，环上几乎没有只差访问顺序就能清的点。"
        "若在覆盖期再按边扫边清专程折返，增量路程将叠加在已经约占虚拟时间 48% 的清除折返之上，覆盖骨干被反复撕开。"
        "相对历史环上边听边清口径 lecture-align（388 s·源），本文先搜后清为 295 s·源，约降 93，故覆盖期应先搜后清。"
    )
    p[56].runs[25].text = ""
    p[56].runs[26].text = ""
    rewrite_plain(
        p[60],
        "表 3 表明：扫完后硬两波、软偏置、改去补测点均抬高秒/源，说明清除折返的主要矛盾不是一次锁定的访问序。"
        "要把城市变成“已经能清的点”，须在环上补全第二次示向；滚动开放路只在此之后才有意义。"
        "环边补测把二测摊进已有环边，顺路清只在增量路程不超过 280 m 时触发，二者都不是边扫边清。",
    )

    insert_after(p[56], "7.2.4  清除路径的滚动开放路优化", style_para=p[46])

    p[57].runs[0].text = "本地对照"
    p[57].runs[1].text = ""
    p[58].runs[0].text = "表 3  清除路径滚动优化本地对照（40 种子，秒·源$^{-1}$）"
    for run in p[58].runs[1:]:
        run.text = ""

    rewrite_plain(p[68], RH_TEXT)

    doc.save(str(OUT))
    replace_fig2(OUT, FIG)
    extract_text(Document(str(OUT)), EXTRACT)
    try:
        copy2(OUT, ORIG)
        copied = True
    except OSError:
        copied = False
    print("saved", OUT)
    print("copied_original", copied)


if __name__ == "__main__":
    main()
