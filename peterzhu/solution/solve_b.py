# -*- coding: utf-8 -*-
"""
2026 高教社杯 B题：无线电干扰源定位与清除
Q1 交会多边形直径 | Q2 第二检测点 | Q3/Q4 本地仿真策略
"""
from __future__ import annotations

from pathlib import Path
import math
import json
import random
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

import numpy as np
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Circle, FancyArrowPatch
from openpyxl import Workbook

OUT = Path(__file__).resolve().parent
FIG = OUT / "figures"
FIG.mkdir(exist_ok=True)

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "SimSun", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False

ARENA_R = 1800.0
EPS_DEG = 1.0
SPEED = 5.0
DET_T = 5.0
SW_T = 1.0
CLEAR_OK_T = 5.0
CLEAR_MISS_T = 3.0
NEAR_R = 5.0
CLEAR_R = 20.0
RX_MIN, RX_MAX = 1000.0, 1500.0


# ========================= 几何工具 =========================
def wrap_deg(a: float) -> float:
    return a % 360.0


def ang_diff_deg(a: float, b: float) -> float:
    """a-b in (-180,180]."""
    d = (a - b + 180.0) % 360.0 - 180.0
    return d


def in_wedge(deg: float, center: float, half: float = EPS_DEG) -> bool:
    return abs(ang_diff_deg(deg, center)) <= half + 1e-9


def bearing_of(p: np.ndarray, g: np.ndarray) -> float:
    v = g - p
    return wrap_deg(math.degrees(math.atan2(v[1], v[0])))


def ray_intersect(p1, ang1_deg, p2, ang2_deg):
    """两射线（半线）交点；平行则 None。射线起点 pi，方向 angi。"""
    a1 = math.radians(ang1_deg)
    a2 = math.radians(ang2_deg)
    d1 = np.array([math.cos(a1), math.sin(a1)])
    d2 = np.array([math.cos(a2), math.sin(a2)])
    A = np.column_stack([d1, -d2])
    if abs(np.linalg.det(A)) < 1e-12:
        return None
    t, s = np.linalg.solve(A, p2 - p1)
    if t < -1e-8 or s < -1e-8:
        return None
    return p1 + t * d1


def point_in_all_wedges(g, stations, bearings, half=EPS_DEG) -> bool:
    for s, b in zip(stations, bearings):
        if not in_wedge(bearing_of(s, g), b, half):
            return False
    return True


def localization_polygon(stations: List[np.ndarray], bearings: List[float], half=EPS_DEG):
    """
    交会定位区域：各检测点误差扇形（±half）的公共交集。
    返回凸多边形顶点（按角排序）。若无界或空，返回 []。
    算法：各站左右边界射线两两求交，筛在所有扇形内的点，再取凸包。
    """
    n = len(stations)
    if n < 2:
        return np.zeros((0, 2))

    cands = []
    # 边界射线角
    left = [wrap_deg(b - half) for b in bearings]
    right = [wrap_deg(b + half) for b in bearings]
    rays = []
    for i in range(n):
        rays.append((stations[i], left[i]))
        rays.append((stations[i], right[i]))

    for i in range(len(rays)):
        for j in range(i + 1, len(rays)):
            # 同站两条射线不相交（于有限点）
            if np.allclose(rays[i][0], rays[j][0]):
                continue
            ipt = ray_intersect(rays[i][0], rays[i][1], rays[j][0], rays[j][1])
            if ipt is None:
                continue
            if point_in_all_wedges(ipt, stations, bearings, half):
                cands.append(ipt)

    if len(cands) < 2:
        return np.zeros((0, 2))

    pts = np.unique(np.round(np.asarray(cands), 6), axis=0)
    if len(pts) < 2:
        return np.zeros((0, 2))

    # 凸包（Andrew）
    pts = pts[np.lexsort((pts[:, 1], pts[:, 0]))]

    def cross(o, a, b):
        return (a[0] - o[0]) * (b[1] - o[1]) - (a[1] - o[1]) * (b[0] - o[0])

    lower = []
    for p in pts:
        while len(lower) >= 2 and cross(lower[-2], lower[-1], p) <= 0:
            lower.pop()
        lower.append(p)
    upper = []
    for p in pts[::-1]:
        while len(upper) >= 2 and cross(upper[-2], upper[-1], p) <= 0:
            upper.pop()
        upper.append(p)
    hull = np.asarray(lower[:-1] + upper[:-1], dtype=float)
    return hull


def polygon_diameter(poly: np.ndarray) -> Tuple[float, Tuple[np.ndarray, np.ndarray]]:
    """凸多边形直径 = 顶点对最大距离。"""
    if len(poly) < 2:
        return 0.0, (None, None)
    best = -1.0
    pair = (poly[0], poly[0])
    m = len(poly)
    for i in range(m):
        for j in range(i + 1, m):
            d = float(np.linalg.norm(poly[i] - poly[j]))
            if d > best:
                best = d
                pair = (poly[i], poly[j])
    return best, pair


def circle_of_diameter_covers(poly: np.ndarray, diam: float, n_test: int = 2000) -> bool:
    """
    以直径线段为直径的圆（半径 diam/2）是否覆盖整个多边形。
    对直径端点取圆周上：任意以 diam 为直径的圆，圆心在直径中点时半径为 diam/2，
    这是所有直径为 diam 的圆中，能覆盖该直径两端点的唯一（Thales）。
    故检验：以最远顶点对中点为圆心、diam/2 为半径，是否覆盖所有顶点。
    """
    if len(poly) == 0 or diam <= 0:
        return True
    _, (a, b) = polygon_diameter(poly)
    c = 0.5 * (a + b)
    r = diam / 2.0 + 1e-6
    for p in poly:
        if np.linalg.norm(p - c) > r:
            return False
    # 额外采样多边形内部
    mn, mx = poly.min(0), poly.max(0)
    for _ in range(n_test):
        q = mn + np.random.rand(2) * (mx - mn)
        if point_in_poly(q, poly) and np.linalg.norm(q - c) > r:
            return False
    return True


def point_in_poly(q, poly):
    # ray casting
    x, y = q
    inside = False
    n = len(poly)
    for i in range(n):
        x1, y1 = poly[i]
        x2, y2 = poly[(i + 1) % n]
        if ((y1 > y) != (y2 > y)) and (x < (x2 - x1) * (y - y1) / (y2 - y1 + 1e-15) + x1):
            inside = not inside
    return inside


# ========================= Q1 =========================
def solve_q1_demo():
    """示例：两站交会 + 三站交会，计算直径并回答圆覆盖问题。"""
    random.seed(0)
    np.random.seed(0)
    G = np.array([400.0, 300.0])
    S1 = np.array([0.0, 0.0])
    S2 = np.array([600.0, 0.0])
    S3 = np.array([200.0, 500.0])

    def noisy_bearing(S, G):
        true = bearing_of(S, G)
        return wrap_deg(true + np.random.uniform(-EPS_DEG, EPS_DEG))

    b1, b2, b3 = noisy_bearing(S1, G), noisy_bearing(S2, G), noisy_bearing(S3, G)

    poly2 = localization_polygon([S1, S2], [b1, b2])
    d2, pair2 = polygon_diameter(poly2)
    cover2 = circle_of_diameter_covers(poly2, d2)

    poly3 = localization_polygon([S1, S2, S3], [b1, b2, b3])
    d3, pair3 = polygon_diameter(poly3)
    cover3 = circle_of_diameter_covers(poly3, d3)

    # 等边三角形反例：直径圆不能覆盖
    eq = np.array([[0.0, 0.0], [1.0, 0.0], [0.5, math.sqrt(3) / 2]])
    deq, _ = polygon_diameter(eq)
    cover_eq = circle_of_diameter_covers(eq, deq)

    fig, axes = plt.subplots(1, 2, figsize=(11, 5), constrained_layout=True)
    for ax, poly, d, pair, title, Ss, bs in [
        (axes[0], poly2, d2, pair2, f"两站交会 直径={d2:.2f}m", [S1, S2], [b1, b2]),
        (axes[1], poly3, d3, pair3, f"三站交会 直径={d3:.2f}m", [S1, S2, S3], [b1, b2, b3]),
    ]:
        if len(poly):
            closed = np.vstack([poly, poly[0]])
            ax.fill(closed[:, 0], closed[:, 1], color="#f4a261", alpha=0.5, label="定位区域")
            ax.plot(closed[:, 0], closed[:, 1], "r-")
            c = 0.5 * (pair[0] + pair[1])
            ax.add_patch(Circle(c, d / 2, fill=False, ls="--", color="blue", label="直径圆"))
            ax.plot([pair[0][0], pair[1][0]], [pair[0][1], pair[1][1]], "b-", lw=1.5)
        ax.scatter([G[0]], [G[1]], c="k", marker="*", s=120, label="真实源", zorder=5)
        for S, b in zip(Ss, bs):
            ax.scatter([S[0]], [S[1]], c="green", s=40)
            rad = math.radians(b)
            ax.arrow(S[0], S[1], 200 * math.cos(rad), 200 * math.sin(rad), head_width=15, color="green", alpha=0.7)
        ax.set_aspect("equal")
        ax.grid(True, alpha=0.3)
        ax.set_title(title)
        ax.legend(fontsize=8)
    fig.savefig(FIG / "Q1_交会定位区域.png", dpi=160)
    plt.close(fig)

    # 等边三角形反例图
    fig, ax = plt.subplots(figsize=(5, 4.5), constrained_layout=True)
    closed = np.vstack([eq, eq[0]])
    ax.fill(closed[:, 0], closed[:, 1], alpha=0.4, color="#e76f51")
    ax.plot(closed[:, 0], closed[:, 1], "k-")
    c = eq.mean(0)  # not diameter center
    a, b = eq[0], eq[1]
    c2 = 0.5 * (a + b)
    ax.add_patch(Circle(c2, deq / 2, fill=False, ls="--", color="blue", label="以底边为直径的圆"))
    ax.scatter([eq[2, 0]], [eq[2, 1]], c="red", s=50, label="未被覆盖的顶点")
    ax.set_aspect("equal")
    ax.legend(fontsize=8)
    ax.set_title("反例：等边三角形不能被直径圆覆盖")
    ax.grid(True, alpha=0.3)
    fig.savefig(FIG / "Q1_直径圆不能覆盖反例.png", dpi=160)
    plt.close(fig)

    result = {
        "two_station_diameter_m": round(d2, 4),
        "two_station_circle_covers": bool(cover2),
        "three_station_diameter_m": round(d3, 4),
        "three_station_circle_covers": bool(cover3),
        "equilateral_triangle_circle_covers": bool(cover_eq),
        "conclusion": "以定位区域直径为直径的圆不一定能覆盖该区域；等边三角形为标准反例（外接圆直径=边长*2/√3>边长）。",
        "algorithm": [
            "各检测点示向度±1°构成误差扇形（楔形）",
            "对所有站左右边界射线两两求交，保留落在全部扇形内的交点",
            "对这些候选点求凸包得到定位多边形",
            "直径=凸包顶点间最大欧氏距离（可用旋转卡壳优化）",
        ],
    }
    (OUT / "Q1_result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Q1:", result["conclusion"])
    print(f"  demo two-station diameter={d2:.3f}, covers={cover2}; triangle covers={cover_eq}")
    return result


# ========================= Q2 =========================
def second_point_candidates(S1: np.ndarray, bearing_deg: float, baseline_range=(300.0, 900.0), n_arc=40):
    """
    第二检测点候选区域：
    - 沿示向度走廊前进，使预期交会角接近 60°–90°
    - 基线长度 L∈[300,900]，落在目标圆内优先
    几何构造：在示向度射线两侧，取与射线夹角约 60°–90° 的点集（环形带）。
    """
    rad = math.radians(bearing_deg)
    u = np.array([math.cos(rad), math.sin(rad)])  # along LOS
    v = np.array([-u[1], u[0]])  # perpendicular

    pts = []
    # 预期源距离未知，在 200–1400 上取若干假设距离 d，再在垂直方向偏置
    for d in np.linspace(200, 1400, 13):
        g_hat = S1 + d * u
        for L in np.linspace(baseline_range[0], baseline_range[1], 8):
            # 使 S1-G-S2 在 G 处近似直角：S2 约在过 G 且垂直于 S1G 的线上
            for sign in (-1, 1):
                # 交会角 φ≈60°~90°：偏置 h = L * sin(alpha)
                for phi in (60, 75, 90):
                    h = L * math.sin(math.radians(phi))
                    # 沿基线方向：从 S1 看，S2 = S1 + L*cos* u_perp混合
                    # 更直接：S2 = g_hat + sign * (L) * rotate
                    ang = math.radians(phi)
                    s2 = g_hat + sign * L * (math.cos(ang) * (-u) + math.sin(ang) * v)
                    # 也加入纯垂直偏置
                    s2b = g_hat + sign * L * v
                    for p in (s2, s2b):
                        pts.append(p)

    pts = np.asarray(pts)
    # 去重
    pts = np.unique(np.round(pts, 1), axis=0)
    # 候选：在圆域稍外也可；标注圆内
    inside = np.linalg.norm(pts, axis=1) <= ARENA_R
    return pts, inside


def score_second_point(S1, b1, S2, true_G=None, rx=1200.0):
    """用交会多边形直径评估（直径越小越好）；无真值时用射线假设距离。"""
    if true_G is not None:
        b2 = bearing_of(S2, true_G)  # 理想无噪声评估几何
    else:
        # 假设源在射线上 800m
        rad = math.radians(b1)
        g_hat = S1 + 800 * np.array([math.cos(rad), math.sin(rad)])
        b2 = bearing_of(S2, g_hat)
    poly = localization_polygon([S1, S2], [b1, b2])
    if len(poly) < 2:
        return 1e9
    d, _ = polygon_diameter(poly)
    # 距离过近或过远惩罚
    baseline = np.linalg.norm(S2 - S1)
    pen = 0.0
    if baseline < 200 or baseline > 1200:
        pen += 50
    return d + pen


def solve_q2_demo():
    S1 = np.array([0.0, 0.0])
    true_G = np.array([700.0, 400.0])
    b1 = wrap_deg(bearing_of(S1, true_G) + 0.3)  # 带误差示向度

    cands, inside = second_point_candidates(S1, b1)
    scores = []
    for p in cands:
        scores.append(score_second_point(S1, b1, p, true_G=true_G))
    scores = np.asarray(scores)
    best_i = int(np.argmin(scores))
    S2_best = cands[best_i]

    # 画候选区与最优
    fig, ax = plt.subplots(figsize=(7, 7), constrained_layout=True)
    ax.add_patch(Circle((0, 0), ARENA_R, fill=False, ls="--", color="gray", label="目标区域"))
    ax.scatter(cands[inside, 0], cands[inside, 1], s=8, c="#90be6d", alpha=0.5, label="候选点(圆内)")
    ax.scatter(cands[~inside, 0], cands[~inside, 1], s=8, c="#adb5bd", alpha=0.3, label="候选点(圆外)")
    ax.scatter([S1[0]], [S1[1]], c="blue", s=60, label="S1")
    ax.scatter([S2_best[0]], [S2_best[1]], c="red", s=80, marker="s", label="推荐 S2")
    ax.scatter([true_G[0]], [true_G[1]], c="k", marker="*", s=140, label="真实源(示意)")
    rad = math.radians(b1)
    ax.arrow(S1[0], S1[1], 900 * math.cos(rad), 900 * math.sin(rad), color="blue", head_width=30, alpha=0.7)
    # 交会多边形
    b2 = bearing_of(S2_best, true_G)
    poly = localization_polygon([S1, S2_best], [b1, b2])
    if len(poly):
        closed = np.vstack([poly, poly[0]])
        ax.fill(closed[:, 0], closed[:, 1], color="#f4a261", alpha=0.55, label="交会区域")
    ax.set_aspect("equal")
    ax.set_xlim(-200, 1600)
    ax.set_ylim(-200, 1600)
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8)
    ax.set_title("问题2：第二检测点候选区域与推荐点")
    fig.savefig(FIG / "Q2_第二检测点候选区.png", dpi=160)
    plt.close(fig)

    # 基线-直径关系曲线
    Ls = np.linspace(100, 1200, 25)
    diams = []
    rad = math.radians(b1)
    u = np.array([math.cos(rad), math.sin(rad)])
    v = np.array([-u[1], u[0]])
    g_hat = true_G
    for L in Ls:
        S2 = g_hat + L * v
        diams.append(score_second_point(S1, b1, S2, true_G=true_G))
    fig, ax = plt.subplots(figsize=(7, 3.8), constrained_layout=True)
    ax.plot(Ls, diams, "o-", color="#264653")
    ax.set_xlabel("基线长度 L / m（垂直偏置）")
    ax.set_ylabel("交会区域直径 / m（越小越好）")
    ax.set_title("问题2：基线长度与定位区域直径")
    ax.grid(True, alpha=0.3)
    fig.savefig(FIG / "Q2_基线与直径.png", dpi=160)
    plt.close(fig)

    result = {
        "strategy": [
            "第一检测点给出示向度走廊（±1°扇形）",
            "第二点宜取较大交会角（约60°–90°），避免近平行射线",
            "推荐基线长度约300–900 m：太短则距离误差放大，太长则可能超出有效接收半径",
            "优先选在示向度射线一侧、使两站视线近似正交的位置",
            "候选区域=沿走廊若干假设距离 × 两侧正交/斜交偏置点集，并优先落在目标圆内",
        ],
        "recommended_S2": [round(float(S2_best[0]), 2), round(float(S2_best[1]), 2)],
        "best_score_diameter": round(float(scores[best_i]), 3),
        "S1": [0, 0],
        "bearing_deg": round(b1, 2),
    }
    (OUT / "Q2_result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Q2 recommended S2", result["recommended_S2"], "diam~", result["best_score_diameter"])
    return result


# ========================= 本地仿真器 + 策略 =========================
@dataclass
class Source:
    channel: int
    pos: np.ndarray
    rx: float
    directional: bool = False
    heading_deg: float = 0.0  # 定向方向：覆盖 heading±90°
    cleared: bool = False

    def covers(self, p: np.ndarray) -> bool:
        if self.cleared:
            return False
        if np.linalg.norm(p - self.pos) > self.rx:
            return False
        if not self.directional:
            return True
        # 从源指向检测点的方位是否在 heading±90
        ang = bearing_of(self.pos, p)
        return abs(ang_diff_deg(ang, self.heading_deg)) <= 90.0 + 1e-9


@dataclass
class LocalSim:
    sources: List[Source]
    pos: np.ndarray = field(default_factory=lambda: np.zeros(2))
    channel: int = 1
    virtual_t: float = 0.0
    n_cleared: int = 0
    seed: int = 0

    def _err(self) -> float:
        # 同点误差固定：用位置哈希
        rng = random.Random(hash((round(self.pos[0], 1), round(self.pos[1], 1), self.seed)) % (2**32))
        return rng.uniform(-EPS_DEG, EPS_DEG)

    def measure(self, x, y, ch: int):
        newp = np.array([x, y], float)
        move = float(np.linalg.norm(newp - self.pos)) / SPEED
        sw = SW_T if ch != self.channel else 0.0
        self.virtual_t += move + sw + DET_T
        self.pos = newp
        self.channel = ch
        src = next((s for s in self.sources if s.channel == ch and not s.cleared), None)
        if src is None or not src.covers(self.pos):
            return {"measure_result": "no_signal", "virtual_time_s": self.virtual_t}
        dist = float(np.linalg.norm(self.pos - src.pos))
        if dist <= NEAR_R and src.covers(self.pos):
            return {"measure_result": "near", "virtual_time_s": self.virtual_t}
        true_b = bearing_of(self.pos, src.pos)
        svd = wrap_deg(true_b + self._err())
        return {"measure_result": "direction", "svd_deg": round(svd, 2), "virtual_time_s": self.virtual_t}

    def clear(self, x, y, ch: int):
        newp = np.array([x, y], float)
        move = float(np.linalg.norm(newp - self.pos)) / SPEED
        src = next((s for s in self.sources if s.channel == ch and not s.cleared), None)
        ok = src is not None and float(np.linalg.norm(newp - src.pos)) <= CLEAR_R
        self.virtual_t += move + (CLEAR_OK_T if ok else CLEAR_MISS_T)
        self.pos = newp
        if ok:
            src.cleared = True
            self.n_cleared += 1
            return {"clear_result": "success", "virtual_time_s": self.virtual_t}
        return {"clear_result": "no_target_in_range", "virtual_time_s": self.virtual_t}


def random_scenario(n_src=None, directional_ratio=0.0, seed=0):
    rng = random.Random(seed)
    if n_src is None:
        n_src = rng.randint(10, 16)
    channels = rng.sample(range(1, 21), n_src)
    sources = []
    n_dir = int(round(n_src * directional_ratio))
    for i, ch in enumerate(channels):
        # 均匀圆盘
        r = ARENA_R * math.sqrt(rng.random())
        ang = rng.random() * 2 * math.pi
        pos = np.array([r * math.cos(ang), r * math.sin(ang)])
        rx = rng.uniform(RX_MIN, RX_MAX)
        directional = i < n_dir
        heading = rng.uniform(0, 360) if directional else 0.0
        sources.append(Source(ch, pos, rx, directional, heading))
    return LocalSim(sources=sources, seed=seed)


def robot_strategy(sim: LocalSim, allow_directional_logic=False, max_steps=800):
    """
    分层策略：
    1) 圆内稀疏侦察点 + 分批扫频发现信号
    2) 两点交会估位后逼近清除
    3) 追击失败则冷却该频道，继续侦察（避免死循环）
    """
    scout = []
    for r in (0, 500, 1000, 1450):
        if r == 0:
            scout.append(np.array([0.0, 0.0]))
        else:
            k = max(6, int(2 * math.pi * r / 800))
            for i in range(k):
                a = 2 * math.pi * i / k + 0.2 * r
                scout.append(np.array([r * math.cos(a), r * math.sin(a)]))

    known = {}  # ch -> list of (pos, svd|'near')
    cleared = set()
    fail_count = {}
    cooldown = set()
    channels = list(range(1, 21))
    steps = 0
    scout_i = 0

    def estimate_from_bearings(stations, bearings):
        if len(stations) >= 2:
            poly = localization_polygon(stations[-2:], bearings[-2:])
            if len(poly):
                return poly.mean(axis=0)
            # 两射线交点（用示向度中心，忽略误差）
            ipt = ray_intersect(stations[-2], bearings[-2], stations[-1], bearings[-1])
            if ipt is not None:
                return ipt
        S, b = stations[-1], bearings[-1]
        rad = math.radians(b)
        return S + 500 * np.array([math.cos(rad), math.sin(rad)])

    def try_localize_and_clear(ch):
        nonlocal steps
        hist = known.get(ch, [])
        if not hist:
            return False
        stations, bearings = [], []
        for p, info in hist:
            if info == "near":
                sim.clear(float(p[0]), float(p[1]), ch)
                steps += 1
                cleared.add(ch)
                return True
            stations.append(np.asarray(p, float))
            bearings.append(float(info))

        if len(stations) == 1:
            S1, b1 = stations[0], bearings[0]
            rad = math.radians(b1)
            # 快速选第二点：沿示向度 600m 处垂直偏置 500m
            g_hat = S1 + 600 * np.array([math.cos(rad), math.sin(rad)])
            v = np.array([-math.sin(rad), math.cos(rad)])
            best_p = g_hat + 500 * v
            if np.linalg.norm(best_p) > ARENA_R:
                best_p = g_hat - 500 * v
            resp = sim.measure(float(best_p[0]), float(best_p[1]), ch)
            steps += 1
            if resp["measure_result"] == "direction":
                known[ch].append((best_p.copy(), resp["svd_deg"]))
                stations.append(best_p)
                bearings.append(resp["svd_deg"])
            elif resp["measure_result"] == "near":
                sim.clear(float(best_p[0]), float(best_p[1]), ch)
                cleared.add(ch)
                return True
            else:
                if allow_directional_logic:
                    alt = g_hat - 500 * v
                    resp2 = sim.measure(float(alt[0]), float(alt[1]), ch)
                    steps += 1
                    if resp2["measure_result"] == "direction":
                        known[ch].append((alt.copy(), resp2["svd_deg"]))
                        stations.append(alt)
                        bearings.append(resp2["svd_deg"])
                    elif resp2["measure_result"] == "near":
                        sim.clear(float(alt[0]), float(alt[1]), ch)
                        cleared.add(ch)
                        return True
                    else:
                        return False
                else:
                    return False

        if len(stations) >= 2:
            guess = estimate_from_bearings(stations, bearings)
            # 沿最后示向度再靠近
            for dist in (0.0, 80.0, 160.0):
                if dist > 0:
                    b = bearings[-1]
                    rad = math.radians(b)
                    tgt = guess + dist * np.array([math.cos(rad), math.sin(rad)])
                else:
                    tgt = guess
                resp = sim.measure(float(tgt[0]), float(tgt[1]), ch)
                steps += 1
                if resp["measure_result"] == "near":
                    sim.clear(float(tgt[0]), float(tgt[1]), ch)
                    cleared.add(ch)
                    return True
                if resp["measure_result"] == "direction":
                    known[ch].append((tgt.copy(), resp["svd_deg"]))
                    rad = math.radians(resp["svd_deg"])
                    for step_d in (40.0, 20.0, 10.0):
                        p2 = tgt + step_d * np.array([math.cos(rad), math.sin(rad)])
                        cr = sim.clear(float(p2[0]), float(p2[1]), ch)
                        steps += 1
                        if cr["clear_result"] == "success":
                            cleared.add(ch)
                            return True
                    # 小范围搜索清除
                    for ang in range(0, 360, 60):
                        a = math.radians(ang)
                        p = tgt + 18 * np.array([math.cos(a), math.sin(a)])
                        cr = sim.clear(float(p[0]), float(p[1]), ch)
                        steps += 1
                        if cr["clear_result"] == "success":
                            cleared.add(ch)
                            return True
                    return False
            return False
        return False

    while steps < max_steps and sim.n_cleared < len(sim.sources):
        pending = [
            ch
            for ch, h in known.items()
            if ch not in cleared and ch not in cooldown and h
        ]
        if pending:
            ch = pending[0]
            ok = try_localize_and_clear(ch)
            if not ok:
                fail_count[ch] = fail_count.get(ch, 0) + 1
                if fail_count[ch] >= 2:
                    cooldown.add(ch)
                    # 保留历史但暂不追击，继续侦察
            continue

        # 冷却重置：侦察一轮后重试
        if scout_i > 0 and scout_i % max(8, len(scout)) == 0:
            cooldown.clear()

        if scout_i >= len(scout):
            a = random.random() * 2 * math.pi
            r = ARENA_R * math.sqrt(random.random())
            scout.append(np.array([r * math.cos(a), r * math.sin(a)]))
        p = scout[scout_i]
        scout_i += 1

        unknown = [c for c in channels if c not in cleared]
        if not unknown:
            break
        start = (scout_i * 3) % len(unknown)
        batch = [unknown[(start + i) % len(unknown)] for i in range(min(5, len(unknown)))]
        for ch in batch:
            resp = sim.measure(float(p[0]), float(p[1]), ch)
            steps += 1
            if resp["measure_result"] == "no_signal":
                continue
            if resp["measure_result"] == "near":
                sim.clear(float(p[0]), float(p[1]), ch)
                cleared.add(ch)
                break
            if resp["measure_result"] == "direction":
                known.setdefault(ch, []).append((p.copy(), resp["svd_deg"]))
                if ch in cooldown:
                    cooldown.discard(ch)
                break

        if sim.virtual_t > 20 * 3600:
            break

    return {
        "cleared": sim.n_cleared,
        "total": len(sim.sources),
        "ratio": sim.n_cleared / max(len(sim.sources), 1),
        "total_time": sim.virtual_t,
        "avg_time": sim.virtual_t / max(sim.n_cleared, 1),
        "steps": steps,
    }


def run_monte_carlo(n_trials=20, directional_ratio=0.0, label="Q3"):
    rows = []
    for seed in range(n_trials):
        sim = random_scenario(directional_ratio=directional_ratio, seed=1000 + seed)
        n_dir = sum(1 for s in sim.sources if s.directional)
        stats = robot_strategy(sim, allow_directional_logic=(directional_ratio > 0))
        stats.update({"seed": seed, "n_dir": n_dir, "label": label})
        rows.append(stats)
        print(
            f"  [{label}] trial {seed}: cleared {stats['cleared']}/{stats['total']} "
            f"ratio={stats['ratio']:.2f} avg_t={stats['avg_time']:.1f}s total_t={stats['total_time']:.1f}s"
        )
    return rows


def plot_scenario_example():
    sim = random_scenario(n_src=12, directional_ratio=0.3, seed=7)
    fig, ax = plt.subplots(figsize=(7, 7), constrained_layout=True)
    ax.add_patch(Circle((0, 0), ARENA_R, fill=False, ls="--", color="gray"))
    for s in sim.sources:
        c = "#e76f51" if s.directional else "#2a9d8f"
        ax.scatter([s.pos[0]], [s.pos[1]], c=c, s=50)
        ax.text(s.pos[0] + 20, s.pos[1] + 20, str(s.channel), fontsize=8)
        if s.directional:
            rad = math.radians(s.heading_deg)
            ax.arrow(s.pos[0], s.pos[1], 120 * math.cos(rad), 120 * math.sin(rad), head_width=25, color=c, alpha=0.6)
            # 覆盖扇形示意
            wedge_theta = np.linspace(s.heading_deg - 90, s.heading_deg + 90, 40)
            wr = 200
            xs = [s.pos[0]] + [s.pos[0] + wr * math.cos(math.radians(t)) for t in wedge_theta] + [s.pos[0]]
            ys = [s.pos[1]] + [s.pos[1] + wr * math.sin(math.radians(t)) for t in wedge_theta] + [s.pos[1]]
            ax.fill(xs, ys, color=c, alpha=0.15)
    ax.scatter([0], [0], c="k", marker="s", s=40, label="起点")
    ax.set_aspect("equal")
    ax.set_title("本地仿真场景示例（绿=全向，红=定向）")
    ax.grid(True, alpha=0.3)
    fig.savefig(FIG / "场景示例_全向与定向.png", dpi=160)
    plt.close(fig)


def write_summary(q3_rows, q4_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "本地仿真统计"
    ws.append(["问题", "试验", "清除数", "总数", "清除比例", "平均定位清除时间s", "总时间s", "定向数"])
    for r in q3_rows + q4_rows:
        ws.append(
            [
                r["label"],
                r["seed"],
                r["cleared"],
                r["total"],
                round(r["ratio"], 4),
                round(r["avg_time"], 2),
                round(r["total_time"], 2),
                r.get("n_dir", 0),
            ]
        )
    # 汇总行
    ws2 = wb.create_sheet("汇总")
    for label, rows in (("Q3全向", q3_rows), ("Q4混合", q4_rows)):
        ratio = np.mean([r["ratio"] for r in rows])
        avg_t = np.mean([r["avg_time"] for r in rows])
        ws2.append([label, "平均清除比例", round(ratio, 4)])
        ws2.append([label, "平均单源时间s", round(avg_t, 2)])
    wb.save(OUT / "B题_本地仿真结果.xlsx")

    # 正式测试表模板（待接模拟器）
    ws3 = Workbook()
    w = ws3.active
    w.title = "表1"
    w.append(["测试案例编码", "清除干扰源个数", "平均定位清除时间", "程序运行时间"])
    w.append(["测试1（待正式测）", "", "", ""])
    w.append(["测试2（待正式测）", "", "", ""])
    w.append(["测试3（待正式测）", "", "", ""])
    ws3.save(OUT / "表1_正式测试模板.xlsx")


def plot_mc(q3_rows, q4_rows):
    fig, ax = plt.subplots(1, 2, figsize=(10, 3.8), constrained_layout=True)
    ax[0].bar(["Q3", "Q4"], [np.mean([r["ratio"] for r in q3_rows]), np.mean([r["ratio"] for r in q4_rows])], color=["#2a9d8f", "#e76f51"])
    ax[0].set_ylim(0, 1.05)
    ax[0].set_ylabel("平均清除比例")
    ax[0].set_title("清除比例")
    ax[0].grid(True, axis="y", alpha=0.3)
    ax[1].bar(
        ["Q3", "Q4"],
        [np.mean([r["avg_time"] for r in q3_rows]), np.mean([r["avg_time"] for r in q4_rows])],
        color=["#2a9d8f", "#e76f51"],
    )
    ax[1].set_ylabel("平均定位清除时间 / s")
    ax[1].set_title("效率")
    ax[1].grid(True, axis="y", alpha=0.3)
    fig.savefig(FIG / "Q3Q4_蒙特卡洛汇总.png", dpi=160)
    plt.close(fig)


# ========================= 官方接口客户端骨架 =========================
CLIENT_CODE = r'''
# robot_client.py — 对接官方模拟器 http://127.0.0.1:2026
# 用法：先启动模拟器并登录，再改 ROBOT_ID 后运行。
import json, time, uuid
from urllib.request import Request, urlopen

BASE_URL = "http://127.0.0.1:2026"
ROBOT_ID = "<参赛队号>"

def post(path, payload):
    req = Request(BASE_URL + path, data=json.dumps(payload).encode("utf-8"),
                  headers={"Content-Type": "application/json"}, method="POST")
    with urlopen(req, timeout=10) as resp:
        return json.loads(resp.read().decode("utf-8"))

def base(rid):
    return {"arena_id": "default", "robot_id": ROBOT_ID, "request_id": rid}

def measure(x, y, ch, rid=None):
    rid = rid or f"m-{uuid.uuid4().hex[:12]}"
    p = base(rid); p["position"]={"x":x,"y":y}; p["channel"]=ch
    return post("/measure", p)

def clear(x, y, ch, rid=None):
    rid = rid or f"c-{uuid.uuid4().hex[:12]}"
    p = base(rid); p["position"]={"x":x,"y":y}; p["channel"]=ch
    return post("/clear", p)

if __name__ == "__main__":
    r = post("/enter", base("enter-1"))
    print(r)
    # 将 solution/solve_b.py 中的 robot_strategy 接到真实 measure/clear 即可
    post("/exit", base("exit-1"))
'''


def main():
    print("=== Q1 ===")
    solve_q1_demo()
    print("=== Q2 ===")
    solve_q2_demo()
    plot_scenario_example()

    print("=== Q3 local MC ===")
    q3_rows = run_monte_carlo(n_trials=8, directional_ratio=0.0, label="Q3")
    print("=== Q4 local MC ===")
    q4_rows = run_monte_carlo(n_trials=8, directional_ratio=0.35, label="Q4")
    write_summary(q3_rows, q4_rows)
    plot_mc(q3_rows, q4_rows)

    (OUT / "robot_client_skeleton.py").write_text(CLIENT_CODE, encoding="utf-8")

    summary = {
        "Q1": json.loads((OUT / "Q1_result.json").read_text(encoding="utf-8")),
        "Q2": json.loads((OUT / "Q2_result.json").read_text(encoding="utf-8")),
        "Q3_mean_ratio": float(np.mean([r["ratio"] for r in q3_rows])),
        "Q3_mean_avg_time_s": float(np.mean([r["avg_time"] for r in q3_rows])),
        "Q4_mean_ratio": float(np.mean([r["ratio"] for r in q4_rows])),
        "Q4_mean_avg_time_s": float(np.mean([r["avg_time"] for r in q4_rows])),
        "note": "Q3/Q4 正式测试需官方模拟器登录；本结果为本地规则一致仿真。",
    }
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print("\nSUMMARY", json.dumps(summary, ensure_ascii=False, indent=2))
    print("outputs ->", OUT)


if __name__ == "__main__":
    main()
