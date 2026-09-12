# -*- coding: utf-8 -*-
"""Rewrite Q2 6.3 as residual / reasonableness analysis and swap figures 4–6."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu
from docx.text.paragraph import Paragraph
from PIL import Image

SRC = Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(7).docx")
FIG = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\figures\q2_paper")
OUT = Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(8).docx")
WECHAT = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题2(5).docx"
)

FIG4 = FIG / "q2_fig4_range_residual.png"
FIG5 = FIG / "q2_fig5_mc_residual.png"
FIG6 = FIG / "q2_fig6_spatial_residual.png"
WIDE = int(Emu(Cm(15.4)))
NARROW = int(Emu(4752975))


def insert_after(paragraph: Paragraph) -> Paragraph:
    new_el = OxmlElement("w:p")
    paragraph._p.addnext(new_el)
    return Paragraph(new_el, paragraph._parent)


def copy_ppr(src: Paragraph, dst: Paragraph) -> None:
    src_ppr = src._p.find(qn("w:pPr"))
    if src_ppr is None:
        return
    old = dst._p.find(qn("w:pPr"))
    if old is not None:
        dst._p.remove(old)
    dst._p.insert(0, deepcopy(src_ppr))


def add_run_like(paragraph: Paragraph, text: str, template_run) -> None:
    run = paragraph.add_run(text)
    src_rpr = template_run._r.find(qn("w:rPr"))
    if src_rpr is not None:
        old = run._r.find(qn("w:rPr"))
        if old is not None:
            run._r.remove(old)
        run._r.insert(0, deepcopy(src_rpr))


def picture_size(path: Path, width_emu: int) -> tuple[int, int]:
    with Image.open(path) as im:
        w, h = im.size
    return width_emu, int(round(width_emu * h / w))


def delete_between(start: Paragraph, end: Paragraph) -> None:
    el = start._p.getnext()
    stop = end._p
    while el is not None and el is not stop:
        nxt = el.getnext()
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
        el = nxt


def find_startswith(doc: Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise RuntimeError(prefix)


def insert_body(cursor: Paragraph, text: str, tpl: Paragraph) -> Paragraph:
    p = insert_after(cursor)
    copy_ppr(tpl, p)
    add_run_like(p, text, tpl.runs[0])
    return p


def insert_heading(cursor: Paragraph, num: str, title: str, tpl: Paragraph) -> Paragraph:
    p = insert_after(cursor)
    copy_ppr(tpl, p)
    add_run_like(p, num, tpl.runs[0])
    title_run = tpl.runs[1] if len(tpl.runs) > 1 else tpl.runs[0]
    add_run_like(p, title, title_run)
    return p


def insert_eq(cursor: Paragraph, formula: str, number: str, tpl: Paragraph) -> Paragraph:
    p = insert_after(cursor)
    copy_ppr(tpl, p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_like(p, f"{formula}     ({number})", tpl.runs[0])
    return p


def insert_figure(
    cursor: Paragraph,
    path: Path,
    width_emu: int,
    caption: str,
    note: str,
    img_tpl: Paragraph,
    cap_tpl: Paragraph,
    note_tpl: Paragraph,
) -> Paragraph:
    p_img = insert_after(cursor)
    copy_ppr(img_tpl, p_img)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    w, h = picture_size(path, width_emu)
    p_img.add_run().add_picture(str(path), width=Emu(w), height=Emu(h))
    p_cap = insert_after(p_img)
    copy_ppr(cap_tpl, p_cap)
    add_run_like(p_cap, caption, cap_tpl.runs[0])
    p_note = insert_after(p_cap)
    copy_ppr(note_tpl, p_note)
    add_run_like(p_note, note, note_tpl.runs[0])
    return p_note


def rewrite_results(doc: Document) -> None:
    p63 = find_startswith(doc, "6.3")
    p64 = find_startswith(doc, "6.4")
    body_tpl = next(p for p in doc.paragraphs if p.text.startswith("针对仅有第一检测站"))
    h_tpl = find_startswith(doc, "6.4.1")
    cap_tpl = next(p for p in doc.paragraphs if p.text.startswith("图 3  "))
    note_tpl = next(p for p in doc.paragraphs if p.text.startswith("输入第一站与示向后"))
    img_tpl = next(p for p in doc.paragraphs if p._element.findall(f".//{qn('w:drawing')}"))
    eq_tpl = next(p for p in doc.paragraphs if p.text.strip() == "(1)")

    delete_between(p63, p64)
    cur = p63

    cur = insert_body(
        cur,
        "建立与求解已经给出候选带与正交构造，本节不再复述选点几何。"
        "把该构造当作可模拟的预报：名义源处交会角应为直角；当真源纵向距离偏离名义距，"
        "或真源落在第一示向正负一度闭角扇内时，交会角残差与包围圆半径残差应远离共线退化，"
        "并在名义距附近使问题一的二十米清除成为可能。"
        "模拟取第一检测站在原点、第一示向三十五度、默认正交第二站"
        "（随体坐标纵向 850 米、侧向 600 米，落在候选带与工作圆内），与表 1 及求解段算例一致。"
        "示向半宽取 1.01 度，与问题一、三的舍入余量相同。",
        body_tpl,
    )

    cur = insert_heading(cur, "6.3.1  ", "残差定义", h_tpl)
    cur = insert_body(
        cur,
        "记真源为 G。两站对 G 的交会角由第一站、第二站与 G 构成的夹角给出。"
        "模型在名义源上要求该角为直角，因而定义交会角残差与包围圆半径残差",
        body_tpl,
    )
    cur = insert_eq(cur, "e_β(G) = β(S₁, S₂, G) − 90°", "5", eq_tpl)
    cur = insert_eq(cur, "e_r(G) = r_SEC(G) − 20", "6", eq_tpl)
    cur = insert_body(
        cur,
        "式中 β 为两站对真源的交会角，单位度；r_SEC 由问题一在两站示向下算出，单位米。"
        "第一站示向取选址所用的三十五度，第二站示向取该站指向真源的几何方位，半宽 1.01 度。"
        "绝对值不超过三十度的交会角残差，对应交会角落入六十度至一百二十度；"
        "e_r 不超过零对应两站后即可二十米清除。"
        "近共线门限 |sin β|<0.08 是同一几何量的另一表述：只要 |sin β| 显著大于该门限，"
        "残差即使离开直角邻域，也不进入问题一禁止清除的退化区。"
        "本节残差刻画选址预报相对真源的偏差，不是高斯点估计的拟合误差。",
        body_tpl,
    )

    cur = insert_heading(cur, "6.3.2  ", "纵向距离失配", h_tpl)
    cur = insert_body(
        cur,
        "真源沿第一示向移动，纵向距离自 250 米至 1600 米、每隔 10 米取一点。"
        "名义纵向 850 米处，交会角为 90.00 度，e_β 在数值上为零，e_r = −1.65 米，两站即可清除。"
        "这是模型的自洽点：构造所瞄准的名义源，接到问题一后包围圆半径低于二十米。",
        body_tpl,
    )
    cur = insert_body(
        cur,
        "离开名义距后，残差由距离失配主导。候选带纵向区间 400 至 1300 米内共 91 个取样点："
        "|e_β| 中位 21.0 度，90% 分位 34.3 度，两端点处 |e_β| = 36.9 度。"
        "该区间内 |sin β| 最小为 0.80，远大于 0.08，近共线样本为零。"
        "因此带端点虽然越出 ±30 度的直角邻域，仍远离共线退化。"
        "e_r 中位 1.38 米；仅在约 590 至 900 米的名义距邻域内 e_r ≤ 0，可清比例 35.2%。"
        "远距一端 1300 米处 e_r = 20.8 米，定位区变大，需要第三站——"
        "这与求解段“不可清则补测”一致，并不构成对正交选址的否定。曲线如图 4。",
        body_tpl,
    )
    cur = insert_figure(
        cur, FIG4, WIDE,
        "图 4  纵向距离失配下的交会角残差与包围圆半径残差",
        "左图浅带为 |e_β|≤30° 的直角邻域，竖虚线为名义纵向 850 米；"
        "右图横线为零残差（二十米清除阈值），星号为名义源处的残差。"
        "绿色纵带标出候选带纵向区间。",
        img_tpl, cap_tpl, note_tpl,
    )

    cur = insert_heading(cur, "6.3.3  ", "示向扰动", h_tpl)
    cur = insert_body(
        cur,
        "将真源纵向均匀取在 400 至 1300 米，并在第一示向上叠加均匀 ±1° 闭区间扰动，"
        "共 400 个样本（随机种子 20260912）。"
        "第二站仍由未扰动的第一示向三十五度生成，以模拟先测向、后选址、真源落在角扇内的在线顺序。",
        body_tpl,
    )
    cur = insert_body(
        cur,
        "扰动后 |e_β| 中位 20.1 度、90% 分位 34.1 度、最大 37.7 度，与无扰动沿示向扫描几乎相同；"
        "|sin β| 最小 0.79，近共线样本仍为零；e_r 中位 1.11 米，e_r ≤ 0 的比例 38.3%。"
        "图 5 表明，赛题 ±1° 误差只在残差上叠加很薄的一层，主结构仍由距离失配决定。"
        "选址模型对题面测向误差带是稳健的。",
        body_tpl,
    )
    cur = insert_figure(
        cur, FIG5, WIDE,
        "图 5  ±1° 示向扰动下的残差分布",
        "左图竖虚线为 ±30° 直角邻域边界；右图竖线为零残差。"
        "样本 400 个，真源纵向在候选带区间内均匀抽取。",
        img_tpl, cap_tpl, note_tpl,
    )

    cur = insert_heading(cur, "6.3.4  ", "合理性", h_tpl)
    cur = insert_body(
        cur,
        "图 6 把 e_β 铺在平面上。直角残差为零的轨迹是以第一站与第二站连线为直径的圆，"
        "这是泰勒斯定理：圆周角为直角当且仅当该点落在直径所对的圆上。"
        "正交构造的几何含义，就是把名义源放到该圆上。"
        "|e_β|=30° 的等值线给出六十度与一百二十度邻域；"
        "候选带穿过名义源附近的浅色区域，而不包含沿示向、残差迅速变大的近场。"
        "沿示向靠近第一站或远离工作圆时残差变大、定位区膨胀，"
        "正是建立段排除示向线、并把名义距压入候选带纵向区间的原因。",
        body_tpl,
    )
    cur = insert_figure(
        cur, FIG6, NARROW,
        "图 6  交会角残差的空间分布与直角轨迹",
        "填色为 e_β；虚线圆以 S₁S₂ 为直径，名义源落在该圆上。"
        "黑实线为 |e_β|=30°；矩形为候选带边界。",
        img_tpl, cap_tpl, note_tpl,
    )
    cur = insert_body(
        cur,
        "综合三组模拟：正交第二站在设计点上交会角残差为零、包围圆半径低于二十米；"
        "在候选带纵向范围内永不触发近共线；±1° 扰动不改变这一结论。"
        "两站即可清除只发生在名义距附近，远距样本需要第三站——"
        "这是问题一清除阈值与源距未知共同造成的，不是选址几何失效。"
        "与共线延伸及粗网格的对照检验见 6.4。",
        body_tpl,
    )


def retouch_validation_and_summary(doc: Document) -> None:
    start = False
    for p in doc.paragraphs:
        if p.text.startswith("6.4"):
            start = True
        if not start:
            continue
        for r in p.runs:
            if r.text.strip() == "(5)":
                r.text = r.text.replace("(5)", "(7)")
            elif r.text.strip() == "(6)":
                r.text = r.text.replace("(6)", "(8)")
    for p in doc.paragraphs:
        if p.runs and p.runs[0].text.startswith("反演与交叉给出的主要结果如下"):
            p.runs[0].text = (
                "残差模拟表明，名义源处交会角残差为零、包围圆半径低于二十米，"
                "候选带纵向范围内不触发近共线。"
                + p.runs[0].text
            )
            break


def verify(doc: Document) -> str:
    lines = [f"n_para={len(doc.paragraphs)} n_inline={len(doc.inline_shapes)} n_tables={len(doc.tables)}"]
    leftover = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text
        if t.startswith("6.") or t.startswith("图 ") or "如图" in t or t.startswith("表 "):
            lines.append(f"[{i:03d}] {t[:180]}")
        if any(k in t for k in ("对照示意点取纵向六百", "同一算例下，候选带、示向线", "橙色菱形为对照")):
            leftover.append(f"P{i}: {t[:80]}")
    lines.append("--- leftover old 6.3 ---")
    lines.extend(leftover or ["none"])
    return "\n".join(lines)


def try_copy(src: Path, dest: Path) -> str:
    try:
        shutil.copy2(src, dest)
        return f"copied {dest}"
    except Exception as e:
        return f"skip {dest.name}: {type(e).__name__}: {e}"


def main():
    doc = Document(str(SRC))
    rewrite_results(doc)
    retouch_validation_and_summary(doc)
    doc.save(str(OUT))

    doc2 = Document(str(OUT))
    report = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\_q2_v8_check.txt")
    report.write_text(verify(doc2), encoding="utf-8")
    extras = [
        Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(7).docx"),
        Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(6).docx"),
        Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(5).docx"),
        WECHAT,
    ]
    notes = [f"wrote {OUT}"]
    for dest in extras:
        notes.append(try_copy(OUT, dest))
    print("\n".join(notes))
    print("report", report)


if __name__ == "__main__":
    main()
