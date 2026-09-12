# -*- coding: utf-8 -*-
"""Revise Q2(5): omni-only 建立/求解, connecting prose, replace 图1/图2."""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import shutil
import tempfile

from docx import Document
from docx.oxml.ns import qn

SRC = Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(5).docx")
FIG = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\figures\q2_paper")
WECHAT = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题2(5).docx"
)
OUT_PROJ = Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(6).docx")

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def set_run(p, idx, text):
    p.runs[idx].text = text


def patch_text(doc: Document) -> None:
    ps = doc.paragraphs

    # 6.1 opening of geometry: drop 「分三步」
    set_run(ps[6], 0, "为把交会质量落实为可计算的选址对象，以第一检测站 ")
    set_run(ps[6], 3, " 由纵向逆时针转九十度。本问默认按全向干扰源处理，左右两侧只要几何上可听，即视为合法候选。")

    set_run(ps[7], 0, "据此沿示向取一名义纵向距离 ")
    set_run(
        ps[7],
        3,
        " 由示向射线与工作圆的交点估计，再截断并压入 400 至 1300 米的候选带纵向区间，"
        "避免名义源过近或过远。至此只确定远近，尚未选定第二站，几何关系如图 1。",
    )

    set_run(ps[9], 0, "图 1  随体坐标、示向轴与名义源")

    set_run(ps[11], 0, "在此基础上，取侧向偏移 ")
    set_run(
        ps[11],
        4,
        " 互相垂直，对名义源的交会角恰为九十度，从而离开共线与反向共线，如图 2（a）。",
    )

    set_run(ps[13], 0, "图 2  侧向正交候选与全向选址流程")

    set_run(
        ps[14],
        3,
        " 为左右侧向候选，阴影为候选带。空心圈标出就近择侧后的推荐站 ",
    )
    set_run(ps[14], 4, "。")

    set_run(ps[15], 0, "若已知当前位置，则在左右合法候选中按增量路程择近，得到推荐第二站 ")
    set_run(
        ps[15],
        1,
        "；越出工作圆则径向缩回，仍不合法则降低侧向高度再试。"
        "图 2（a）中空心圈即择近结果：默认档两侧到第一站路程相近时，再取距工作圆圆心更近者。"
        "全向情形下左右候选对称可听，不再附加朝向半空间约束。",
    )

    set_run(
        ps[16],
        0,
        "与上述正交构造相应，实现中可对照三档参数，正文只以默认正交档为展开对象。"
        "对照示意点取纵向约六百米、侧向正负三百五十米，并要求落在三圆盘可听交内，侧向低于候选带下界，仅作对照。"
        "紧凑正交取纵向约四百五十米、侧向正负四百米，路程更短但侧向裕度较小。默认正交取侧向",
    )

    set_run(
        ps[17],
        8,
        "米；角点出圆则径向缩回。侧向为零、沿示向前进的点不在候选区内，因其使交会角趋于共线、"
        "定位区细长不可清，已由问题一给出。可听性用三圆盘交作辅助：以第一检测站为心，以及沿示向误差上下界射线各取一千米处为心，"
        "半径均为一千米；对照档要求第二站落在三者之交内。该保证不意味着两站之后包围圆半径必不超过二十米。"
        "本问默认全向干扰源，故不引入未知朝向的前半空间约束。",
    )

    set_run(
        ps[21],
        0,
        "仅有一站示向时，全平面网格搜索最坏交会直径代价高、也不便在线解释。"
        "本问按全向干扰源作构造性正交选址：在随体坐标中取名义纵向距离，向侧向生成左右候选，"
        "经候选带与工作圆裁剪、合法性筛选与就近择侧后输出第二站，再回代问题一。"
        "计算流程如图 2（b）。网格搜索只作交叉验证对照，不写入主求解。",
    )

    # 6.2.3 drop directional half-space; last oMath becomes unused
    set_run(
        ps[33],
        3,
        " 表示取使后面目标达到最小的自变量：这里是在左右合法候选中比较两点到当前位置的距离，"
        "返回距离更短的那一个点，而不是返回最小距离数值本身。若两路程相等，再取距工作圆圆心更近者。"
        "本问按全向源处理，左右合法点均保留，只按路程择侧。",
    )
    set_run(ps[33], 4, "")
    oms = ps[33]._element.findall(f".//{{{M}}}oMath")
    if len(oms) >= 4:
        parent = oms[-1].getparent()
        if parent is not None:
            parent.remove(oms[-1])

    set_run(
        ps[51],
        0,
        "本问验证针对几何约束满足与接到问题一后的可清性自洽，不以预测误差为主判据，"
        "亦不以第二站越近越好为标准（过近会使交会细长，问题一已拒绝近共线）。验证分为反演与交叉两类。",
    )
    set_run(
        ps[55],
        2,
        "。名义源旁侧的正交点满足上式；对照示意点侧向低于带下界时不作默认推荐。",
    )

    set_run(
        ps[71],
        0,
        "求解过程与前述各小节一致。先建立随体坐标，由示向与工作圆求交得到名义纵向距离并压入候选带；"
        "再按正交构造生成左右候选，默认侧向取满；随后做合法性筛选并就近择侧。"
        "第二站测向后回代问题一，若仍不可清，则触发第三站补测。"
        "本问默认全向干扰源，不附加朝向半空间约束。"
        "算例中默认推荐站随体坐标约纵向",
    )


def replace_media(docx_path: Path) -> None:
    img1 = (FIG / "q2_fig1_nominal.png").read_bytes()
    img2 = (FIG / "q2_fig2_ortho_flow.png").read_bytes()
    mapping = {
        "word/media/image1.png": img1,
        "word/media/image2.png": img2,
        "word/media/image8.png": img1,
        "word/media/image9.png": img2,
    }
    tmp = docx_path.with_suffix(".tmp.docx")
    with ZipFile(docx_path, "r") as zin, ZipFile(tmp, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = mapping.get(item.filename, zin.read(item.filename))
            zout.writestr(item, data)
    tmp.replace(docx_path)


def main():
    shutil.copy2(SRC, OUT_PROJ)
    doc = Document(str(OUT_PROJ))
    patch_text(doc)
    doc.save(str(OUT_PROJ))
    replace_media(OUT_PROJ)

    # verify
    doc2 = Document(str(OUT_PROJ))
    flags = []
    for i, p in enumerate(doc2.paragraphs):
        t = p.text
        if any(k in t for k in ("第一步", "第二步", "第三步", "定向", "问题四预留", "未知朝向", "前半空间", "背面")):
            if i < 72:
                flags.append(f"P{i}: {t[:80]}")
    report = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\_q2_v6_check.txt")
    lines = [f"imgs={len(doc2.inline_shapes)}"]
    for i in (6, 7, 9, 11, 13, 14, 15, 17, 21, 33, 51, 55, 71):
        lines.append(f"[{i}] {doc2.paragraphs[i].text[:220]}")
    lines.append("--- leftover ---")
    lines.extend(flags or ["none"])
    report.write_text("\n".join(lines), encoding="utf-8")

    for dest in (SRC, WECHAT):
        try:
            shutil.copy2(OUT_PROJ, dest)
            print("copied", dest)
        except Exception as e:
            print("skip", dest.name, type(e).__name__)
    print("wrote", OUT_PROJ)


if __name__ == "__main__":
    main()
