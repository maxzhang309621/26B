# -*- coding: utf-8 -*-
"""Split Q2 fig2: keep ortho schematic, move flow to a new fig3 in 6.2."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu
from docx.text.paragraph import Paragraph
from PIL import Image

SRC = Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(6).docx")
FIG = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\figures\q2_paper")
OUT = Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(7).docx")
WECHAT = Path(
    r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
    r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题2(5).docx"
)
FIG2 = FIG / "q2_fig2_ortho.png"
FIG3 = FIG / "q2_fig3_flow.png"

FIG1_WIDTH = Emu(4752975)  # keep 图1 column width


def insert_after(paragraph: Paragraph) -> Paragraph:
    new_el = OxmlElement("w:p")
    paragraph._p.addnext(new_el)
    return Paragraph(new_el, paragraph._parent)


def copy_ppr(src: Paragraph, dst: Paragraph) -> None:
    src_ppr = src._p.find(qn("w:pPr"))
    if src_ppr is None:
        return
    old = dst._p.find(qn("w:pPr"))
    if old is not None:
        dst._p.remove(old)
    dst._p.insert(0, deepcopy(src_ppr))


def add_run_like(paragraph: Paragraph, text: str, template_run) -> None:
    run = paragraph.add_run(text)
    src_rpr = template_run._r.find(qn("w:rPr"))
    if src_rpr is not None:
        old = run._r.find(qn("w:rPr"))
        if old is not None:
            run._r.remove(old)
        run._r.insert(0, deepcopy(src_rpr))


def clear_runs(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def picture_size(path: Path, width_emu: int) -> tuple[int, int]:
    with Image.open(path) as im:
        w, h = im.size
    height_emu = int(round(width_emu * h / w))
    return width_emu, height_emu


def set_run_if(p, idx, text):
    p.runs[idx].text = text


def patch_text(doc: Document) -> None:
    ps = doc.paragraphs
    set_run_if(
        ps[11],
        4,
        " 互相垂直，对名义源的交会角恰为九十度，从而离开共线与反向共线，如图 2。",
    )
    set_run_if(ps[13], 0, "图 2  侧向正交候选与就近推荐站")
    set_run_if(
        ps[15],
        1,
        "；越出工作圆则径向缩回，仍不合法则降低侧向高度再试。"
        "图 2 中空心圈即择近结果：默认档两侧到第一站路程相近时，再取距工作圆圆心更近者。"
        "全向情形下左右候选对称可听。",
    )
    set_run_if(
        ps[21],
        0,
        "仅有一站示向时，全平面网格搜索最坏交会直径代价高、也不便在线解释。"
        "本问按全向干扰源作构造性正交选址：在随体坐标中取名义纵向距离，向侧向生成左右候选，"
        "经候选带与工作圆裁剪、合法性筛选与就近择侧后输出第二站，再回代问题一。"
        "计算流程如图 3。",
    )

    # 6.3 起旧图 3–7 顺延为 4–8，从大到小以免误伤
    start = next(i for i, p in enumerate(ps) if p.text.startswith("6.3"))
    for old, new in ((7, 8), (6, 7), (5, 6), (4, 5), (3, 4)):
        token_old = f"图 {old}"
        token_new = f"图 {new}"
        spaced_old = f" {old} "
        spaced_new = f" {new} "
        for p in ps[start:]:
            for r in p.runs:
                if token_old in r.text:
                    r.text = r.text.replace(token_old, token_new)
                elif r.text == spaced_old:
                    r.text = spaced_new
                elif r.text == f" {old}":
                    # 「（图 5）」拆成 r3=' 5' r4='）'
                    r.text = f" {new}"


def replace_fig2(doc: Document) -> None:
    p = doc.paragraphs[12]
    clear_runs(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    w, h = picture_size(FIG2, int(FIG1_WIDTH))
    run = p.add_run()
    run.add_picture(str(FIG2), width=Emu(w), height=Emu(h))


def insert_fig3(doc: Document) -> None:
    ps = doc.paragraphs
    p21 = ps[21]
    p_img_tpl = ps[12]
    p_cap_tpl = ps[13]
    p_note_tpl = ps[14]

    p_img = insert_after(p21)
    copy_ppr(p_img_tpl, p_img)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 流程图更扁，略加宽
    w, h = picture_size(FIG3, int(Emu(Cm(15.4))))
    run = p_img.add_run()
    run.add_picture(str(FIG3), width=Emu(w), height=Emu(h))

    p_cap = insert_after(p_img)
    copy_ppr(p_cap_tpl, p_cap)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_like(p_cap, "图 3  全向第二站选址流程", p_cap_tpl.runs[0])

    p_note = insert_after(p_cap)
    copy_ppr(p_note_tpl, p_note)
    add_run_like(
        p_note,
        "输入第一站与示向后，由射线与工作圆求交得到名义纵向距离，再生成左右正交候选并裁剪到工作圆；"
        "经候选带筛选与就近择侧得到第二站。若回代问题一后包围圆半径不超过二十米则可清，否则补第三站。",
        p_note_tpl.runs[0],
    )

    p_grid = insert_after(p_note)
    copy_ppr(p21, p_grid)
    add_run_like(
        p_grid,
        "网格搜索只作交叉验证对照，不写入主求解。",
        p21.runs[0],
    )


def verify(doc: Document) -> str:
    lines = [f"n_para={len(doc.paragraphs)} n_inline={len(doc.inline_shapes)}"]
    hits = []
    leftover = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text
        if any(k in t for k in ("图 2（a）", "图 2（b）", "如图 2（", "图2（")):
            leftover.append(f"P{i}: {t[:120]}")
        if t.startswith("图 ") or "如图" in t or t.startswith("6."):
            hits.append(f"[{i:03d}] {t[:160]}")
    lines.append("--- structure ---")
    lines.extend(hits)
    lines.append("--- leftover a/b ---")
    lines.extend(leftover or ["none"])
    return "\n".join(lines)


def try_copy(src: Path, dest: Path) -> str:
    try:
        shutil.copy2(src, dest)
        return f"copied {dest}"
    except Exception as e:
        return f"skip {dest.name}: {type(e).__name__}: {e}"


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    patch_text(doc)
    replace_fig2(doc)
    insert_fig3(doc)
    doc.save(str(OUT))

    doc2 = Document(str(OUT))
    report = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\_q2_v7_check.txt")
    report.write_text(verify(doc2), encoding="utf-8")

    extras = [
        SRC,
        Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(5).docx"),
        WECHAT,
    ]
    notes = [f"wrote {OUT}"]
    for dest in extras:
        notes.append(try_copy(OUT, dest))
    print("\n".join(notes))
    print("report", report)


if __name__ == "__main__":
    main()
