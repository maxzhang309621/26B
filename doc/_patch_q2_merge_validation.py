# -*- coding: utf-8 -*-
"""Drop 结果分析; merge residual into 模型验证; replace inversion; remove cross-validation."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu
from docx.text.paragraph import Paragraph
from PIL import Image

SRC = Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(8).docx")
OUT = Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(9).docx")
FIG7 = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\figures\q2_paper\q2_fig7_inversion.png")
WIDE = int(Emu(Cm(15.4)))


def insert_after(paragraph: Paragraph) -> Paragraph:
    new_el = OxmlElement("w:p")
    paragraph._p.addnext(new_el)
    return Paragraph(new_el, paragraph._parent)


def copy_ppr(src: Paragraph, dst: Paragraph) -> None:
    src_ppr = src._p.find(qn("w:pPr"))
    if src_ppr is None:
        return
    old = dst._p.find(qn("w:pPr"))
    if old is not None:
        dst._p.remove(old)
    dst._p.insert(0, deepcopy(src_ppr))


def add_run_like(paragraph: Paragraph, text: str, template_run) -> None:
    run = paragraph.add_run(text)
    src_rpr = template_run._r.find(qn("w:rPr"))
    if src_rpr is not None:
        old = run._r.find(qn("w:rPr"))
        if old is not None:
            run._r.remove(old)
        run._r.insert(0, deepcopy(src_rpr))


def picture_size(path: Path, width_emu: int) -> tuple[int, int]:
    with Image.open(path) as im:
        w, h = im.size
    return width_emu, int(round(width_emu * h / w))


def delete_between(start: Paragraph, end: Paragraph) -> None:
    el = start._p.getnext()
    stop = end._p
    while el is not None and el is not stop:
        nxt = el.getnext()
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)
        el = nxt


def find_startswith(doc: Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p
    raise RuntimeError(prefix)


def find_contains(doc: Document, token: str) -> Paragraph:
    for p in doc.paragraphs:
        if token in p.text:
            return p
    raise RuntimeError(token)


def set_first_run_containing(p: Paragraph, old: str, new: str) -> None:
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return
    raise RuntimeError(old)


def insert_body(cursor: Paragraph, text: str, tpl: Paragraph) -> Paragraph:
    p = insert_after(cursor)
    copy_ppr(tpl, p)
    add_run_like(p, text, tpl.runs[0])
    return p


def insert_heading(cursor: Paragraph, num: str, title: str, tpl: Paragraph) -> Paragraph:
    p = insert_after(cursor)
    copy_ppr(tpl, p)
    add_run_like(p, num, tpl.runs[0])
    title_run = tpl.runs[1] if len(tpl.runs) > 1 else tpl.runs[0]
    add_run_like(p, title, title_run)
    return p


def insert_figure(cursor, path, width_emu, caption, note, img_tpl, cap_tpl, note_tpl):
    p_img = insert_after(cursor)
    copy_ppr(img_tpl, p_img)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    w, h = picture_size(path, width_emu)
    p_img.add_run().add_picture(str(path), width=Emu(w), height=Emu(h))
    p_cap = insert_after(p_img)
    copy_ppr(cap_tpl, p_cap)
    add_run_like(p_cap, caption, cap_tpl.runs[0])
    p_note = insert_after(p_cap)
    copy_ppr(note_tpl, p_note)
    add_run_like(p_note, note, note_tpl.runs[0])
    return p_note


def insert_table_after(doc: Document, paragraph: Paragraph, rows: list[list[str]]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    if doc.tables:
        try:
            table.style = doc.tables[0].style
        except Exception:
            pass
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            table.cell(i, j).text = val
    tbl = table._tbl
    parent = tbl.getparent()
    if parent is not None:
        parent.remove(tbl)
    paragraph._p.addnext(tbl)
    return table


def patch(doc: Document) -> None:
    body_tpl = next(p for p in doc.paragraphs if p.text.startswith("针对仅有第一检测站"))
    h_tpl = find_startswith(doc, "6.3.1")
    cap_tpl = next(p for p in doc.paragraphs if p.text.startswith("图 4  "))
    note_tpl = next(p for p in doc.paragraphs if p.text.startswith("左图浅带"))
    img_tpl = next(p for p in doc.paragraphs if p._element.findall(f".//{qn('w:drawing')}"))

    p63 = find_startswith(doc, "6.3")
    set_first_run_containing(p63, "结果分析", "模型验证")

    intro = find_startswith(doc, "建立与求解已经给出候选带")
    intro.runs[0].text = (
        "本问验证针对选址预报的残差与反演自洽，不以坐标预测误差为主判据，亦不以第二站越近越好为标准。"
        "先把正交构造当作可模拟的预报：名义源处交会角应为直角；当真源纵向距离偏离名义距，"
        "或真源落在第一示向正负一度闭角扇内时，交会角残差与包围圆半径残差应远离共线退化。"
        "再把两站示向送回问题一，检验真源是否落入角扇交、示向残差是否落在误差带内，并与共线延伸站对照。"
        "模拟取第一检测站在原点、第一示向三十五度、默认正交第二站"
        "（随体坐标纵向 850 米、侧向 600 米，落在候选带与工作圆内），与表 1 及求解段算例一致。"
        "示向半宽取 1.01 度，与问题一、三的舍入余量相同。"
    )
    if len(intro.runs) > 1:
        for r in intro.runs[1:]:
            r.text = ""

    close = find_contains(doc, "与共线延伸及粗网格的对照检验见")
    set_first_run_containing(
        close,
        "与共线延伸及粗网格的对照检验见 6.4。",
        "测量反演与共线对照见下一小节。",
    )

    p_sum = find_startswith(doc, "6.5")
    delete_between(close, p_sum)

    cur = close
    cur = insert_heading(cur, "6.3.5  ", "反演验证", h_tpl)
    cur = insert_body(
        cur,
        "几何反演核对候选带与正交构造。默认正交站随体坐标为纵向 850 米、侧向 600 米，落入式（3）的候选带且在工作圆内，"
        "对名义源交会角为 90.00 度。沿示向前进 800 米、侧向为零的点不在候选带内，对应题意对示向线上点的排除。"
        "对照讲义点侧向 350 米低于带下界 400 米，亦不在带内。上述三点与单元测试一致，不依赖随机抽样。",
        body_tpl,
    )
    cur = insert_body(
        cur,
        "测量反演把第二站接到问题一。真源纵向均匀取在 400 至 1300 米，相对第一示向叠加均匀 ±1° 闭区间扰动；"
        "第二站示向再叠加独立 ±1° 扰动，并舍入到两位小数，共 400 个样本（随机种子 20260912）。"
        "第一站示向固定为选址所用的三十五度。两方案共用同一批真源：一为正交推荐站，一为沿示向 800 米的共线延伸站。"
        "对尚未可清的样本，沿定位区长轴法向补第三站后再测一次。",
        body_tpl,
    )
    cur = insert_body(
        cur,
        "正交与共线的真值包含率、示向残差落在 ±1.01° 带内的比例均为 100%，正交方案示向残差最大绝对值 0.998 度。"
        "这说明问题一的角扇交会核与测量生成规则一致，两种选址都能“反演出”真源所在的角扇。"
        "差别出现在交会质量。正交方案对真源的交会角中位 86.7 度，近共线 0，两站可清 33.2%，"
        "包围圆半径中位 21.6 米；对尚未可清样本补第三站后，本批 100% 可清。"
        "共线方案交会角中位 4.9 度，近共线 81.2%，两站可清仅 4.0%；补第三站后可清 65.5%，"
        "且有限包围圆半径中位仍达 47.9 米。"
        "据此，第二站目标是避免共线退化、为后续测站留下可收缩定位区，而非承诺任意源距下两站必有二十米清除。"
        "对照如图 7 与表 2。",
        body_tpl,
    )
    cur = insert_figure(
        cur, FIG7, WIDE,
        "图 7  反演验证：正交推荐站与共线延伸站对照",
        "同一批 400 个真源。包含率与残差带通过率检验问题一的角扇核；"
        "可清与近共线检验选址是否离开退化。第三站沿定位区长轴法向补测。",
        img_tpl, cap_tpl, note_tpl,
    )
    p_tab_cap = insert_after(cur)
    copy_ppr(cap_tpl, p_tab_cap)
    add_run_like(p_tab_cap, "表 2  反演对照：正交推荐与共线延伸", cap_tpl.runs[0])
    insert_table_after(
        doc,
        p_tab_cap,
        [
            ["方案", "真值包含", "两站可清", "近共线", "补第三站后可清", "交会角中位"],
            ["正交推荐", "100%", "33.2%", "0", "100%", "86.7°"],
            ["共线延伸", "100%", "4.0%", "81.2%", "65.5%", "4.9°"],
        ],
    )

    for p in doc.paragraphs:
        for r in p.runs:
            if "网格搜索只作交叉验证对照，不写入主求解" in r.text:
                r.text = r.text.replace(
                    "网格搜索只作交叉验证对照，不写入主求解",
                    "网格搜索最坏直径不写入主求解",
                )
            if "网格搜索最坏直径只作为交叉验证对照，不写入主求解" in r.text:
                r.text = r.text.replace(
                    "网格搜索最坏直径只作为交叉验证对照，不写入主求解",
                    "网格搜索最坏直径不写入主求解",
                )

    p_sum = find_startswith(doc, "6.5")
    set_first_run_containing(p_sum, "6.5", "6.4")

    tail = find_startswith(doc, "残差模拟表明")
    tail.runs[0].text = (
        "残差模拟表明，名义源处交会角残差为零、包围圆半径低于二十米，候选带纵向范围内不触发近共线；"
        "±1° 示向扰动不改变这一结构。测量反演中，正交与共线都能把真源包含在角扇交内、示向残差落在 ±1.01° 带内，"
        "但正交方案交会角中位 86.7 度、近共线为 0、两站可清 33.2%，补第三站后本批全部可清；"
        "共线延伸交会角中位 4.9 度、近共线 81.2%、两站可清 4.0%，补第三站后仍有约三分之一不可清。"
        "正交相对共线的差别，在于能否为问题一留下可收缩、可清的定位区。"
        "本问最终向问题三、四交付默认正交第二站选址接口；后问在此基础上组织巡航与清除，不再另写第二站选址过程。"
    )
    if len(tail.runs) > 1:
        for r in tail.runs[1:]:
            r.text = ""


def verify(doc: Document) -> str:
    lines = [f"n_para={len(doc.paragraphs)} n_inline={len(doc.inline_shapes)} n_tables={len(doc.tables)}"]
    bad = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text
        if t.startswith("6.") or t.startswith("图 ") or t.startswith("表 ") or "如图" in t:
            lines.append(f"[{i:03d}] {t[:200]}")
        if any(k in t for k in ("结果分析", "交叉验证", "6.4.2", "6.5  ", "图 8", "51.7", "99 组")):
            bad.append(f"P{i}: {t[:100]}")
    lines.append("--- leftover ---")
    lines.extend(bad or ["none"])
    return "\n".join(lines)


def try_copy(src: Path, dest: Path) -> str:
    try:
        shutil.copy2(src, dest)
        return f"copied {dest.name}"
    except Exception as e:
        return f"skip {dest.name}: {type(e).__name__}"


def main():
    doc = Document(str(SRC))
    patch(doc)
    doc.save(str(OUT))
    doc2 = Document(str(OUT))
    report = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\_q2_v9_check.txt")
    report.write_text(verify(doc2), encoding="utf-8")
    extras = [
        Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(8).docx"),
        Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(6).docx"),
        Path(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(5).docx"),
        Path(
            r"c:\Users\max\xwechat_files\wxid_whiruqlnhzi712_2e9a\temp\RWTemp"
            r"\2026-09\805970c12c02fdc51e5af8cbfb1bb35a\2026B-问题2(5).docx"
        ),
    ]
    print("wrote", OUT)
    for dest in extras:
        print(try_copy(OUT, dest))
    print("report", report)


if __name__ == "__main__":
    main()
