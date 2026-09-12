# -*- coding: utf-8 -*-
from docx import Document
from docx.oxml.ns import qn
from pathlib import Path

doc = Document(r"d:\maxzhang\python files\Math-modeling\26B\doc\2026B-问题2(6).docx")
out = Path(r"d:\maxzhang\python files\Math-modeling\26B\output\_q2_inspect.txt")
lines = []
lines.append(f"n_para={len(doc.paragraphs)} n_inline={len(doc.inline_shapes)}")
for k, s in enumerate(doc.inline_shapes):
    lines.append(f"shape{k} w={s.width} h={s.height}")
for i in range(len(doc.paragraphs)):
    p = doc.paragraphs[i]
    drawings = p._element.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing")
    lines.append(f"[{i:03d}] align={p.alignment} runs={len(p.runs)} drawings={len(drawings)}")
    lines.append("    TEXT: " + p.text.replace("\n", " | "))
    for j, r in enumerate(p.runs):
        if r.text:
            lines.append(f"    r{j}: {r.text}")
out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out)
