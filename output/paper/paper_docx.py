# -*- coding: utf-8 -*-
"""CUMCM Word helpers aligned with cumcm-2026b-paper format-rules."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

SRC = Path(__file__).resolve().parents[2] / "src"
sys.path.insert(0, str(SRC))
from word_math import (  # noqa: E402
    add_equation,
    add_mixed_paragraph,
    _append_text_runs,
    _set_no_midword_wrap,
)

BODY_SIZE = 10.5
H1_SIZE = 15.0
H2_SIZE = 12.0
H3_SIZE = 10.5
CAPTION_SIZE = 9.0
REF_SIZE = 9.0
BODY_LINE_SPACING = 1.0
HEADING_SPACE_AFTER = Pt(5)
CITE_RE = re.compile(r"\[(\d+)\]")
FONT_LATIN = "Times New Roman"
FONT_BODY_CN = "宋体"
FONT_HEAD_CN = "黑体"


def set_run_font(
    run,
    size=BODY_SIZE,
    bold=False,
    superscript=False,
    *,
    latin=FONT_LATIN,
    east_asia=FONT_BODY_CN,
):
    """Skill: 正文/表 — 中文 eastAsia=宋体，拉丁 ascii/hAnsi=Times New Roman。"""
    run.bold = bold
    run.font.size = Pt(size)
    run.font.superscript = superscript
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "cs"):
        rFonts.set(qn(f"w:{key}"), latin)
    rFonts.set(qn("w:eastAsia"), east_asia)


def set_heading_font(run, size, bold=True):
    set_run_font(run, size, bold, east_asia=FONT_HEAD_CN)


def _para_spacing(pf, *, first_indent=True, followup=False):
    pf.line_spacing = BODY_LINE_SPACING
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if followup:
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
    else:
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        if first_indent:
            pf.first_line_indent = Cm(0.74)


def _set_cell_border(cell, **edges):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tc_borders = tcPr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tcPr.append(tc_borders)
    for edge, spec in edges.items():
        tag = qn(f"w:{edge}")
        el = tc_borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_borders.append(el)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), spec.get("val", "single"))
            el.set(qn("w:sz"), str(spec.get("sz", 4)))
            el.set(qn("w:color"), spec.get("color", "auto"))


def _apply_three_line_table(table):
    rows = table.rows
    n = len(rows)
    nil = {"val": "nil"}
    top = {"val": "single", "sz": 12}
    mid = {"val": "single", "sz": 6}
    bot = {"val": "single", "sz": 12}
    for i, row in enumerate(rows):
        for cell in row.cells:
            if i == 0:
                _set_cell_border(cell, top=top, bottom=mid, left=nil, right=nil)
            elif i == n - 1:
                _set_cell_border(cell, bottom=bot, left=nil, right=nil)
            else:
                _set_cell_border(cell, left=nil, right=nil)


class PaperDoc:
    def __init__(self, title: str | None = None):
        self.doc = Document()
        self.fig_no = 0
        self.tbl_no = 0
        self.eq_no = 0
        self.cites_used: set[int] = set()
        for sec in self.doc.sections:
            sec.top_margin = Cm(2.54)
            sec.bottom_margin = Cm(2.54)
            sec.left_margin = Cm(2.5)
            sec.right_margin = Cm(2.5)
        if title:
            self.add_title(title)
        self._disable_hyphenation()

    def _disable_hyphenation(self):
        settings = self.doc.settings.element
        for tag, val in (("autoHyphenation", "0"), ("doNotHyphenateCaps", "1")):
            el = settings.find(qn(f"w:{tag}"))
            if el is None:
                el = OxmlElement(f"w:{tag}")
                settings.append(el)
            el.set(qn("w:val"), val)

    def save(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        print(f"Wrote {path}")

    def add_title(self, text: str):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_heading_font(run, 16)

    def add_heading1(self, text: str):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_no_midword_wrap(p)
        _append_text_runs(p, text, size=H1_SIZE, bold=True, east_asia=FONT_HEAD_CN)
        p.paragraph_format.space_after = HEADING_SPACE_AFTER

    def add_heading2(self, text: str):
        p = self.doc.add_paragraph()
        _set_no_midword_wrap(p)
        _append_text_runs(p, text, size=H2_SIZE, bold=True, east_asia=FONT_HEAD_CN)
        p.paragraph_format.space_after = HEADING_SPACE_AFTER

    def add_heading3(self, text: str):
        p = self.doc.add_paragraph()
        _set_no_midword_wrap(p)
        _append_text_runs(p, text, size=H3_SIZE, bold=True, east_asia=FONT_HEAD_CN)
        p.paragraph_format.space_after = HEADING_SPACE_AFTER

    def _append_runs(self, p, text: str, *, size=BODY_SIZE, bold=False):
        _set_no_midword_wrap(p)
        pos = 0
        for m in CITE_RE.finditer(text):
            if m.start() > pos:
                _append_text_runs(p, text[pos : m.start()], size=size, bold=bold)
            num = int(m.group(1))
            self.cites_used.add(num)
            run = p.add_run(f"[{num}]")
            set_run_font(run, BODY_SIZE, False, superscript=True)
            pos = m.end()
        if pos < len(text):
            _append_text_runs(p, text[pos:], size=size, bold=bold)

    def add_para(self, text: str, *, first_indent=True):
        if "$" in text:
            p = add_mixed_paragraph(
                self.doc,
                text,
                first_indent=first_indent,
                font_size=BODY_SIZE,
                line_spacing=BODY_LINE_SPACING,
            )
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.line_spacing = BODY_LINE_SPACING
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            for m in CITE_RE.finditer(text):
                self.cites_used.add(int(m.group(1)))
            return p
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _para_spacing(p.paragraph_format, first_indent=first_indent)
        _set_no_midword_wrap(p)
        self._append_runs(p, text)
        return p

    def add_eq_followup(self, text: str):
        if "$" in text:
            p = add_mixed_paragraph(
                self.doc,
                text,
                first_indent=False,
                font_size=BODY_SIZE,
                line_spacing=BODY_LINE_SPACING,
            )
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.line_spacing = BODY_LINE_SPACING
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.first_line_indent = Cm(0)
            pf.space_before = Pt(0)
            return p
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _para_spacing(p.paragraph_format, first_indent=False, followup=True)
        _set_no_midword_wrap(p)
        self._append_runs(p, text)
        return p

    def add_display_eq(self, latex: str, followup: str | None = None):
        self.eq_no += 1
        add_equation(self.doc, latex, numbered=f"({self.eq_no})", font_size=BODY_SIZE)
        if followup:
            self.add_eq_followup(followup)

    def _fill_cell(self, cell, text: str, *, bold: bool = False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if "$" in text:
            add_mixed_paragraph(
                self.doc,
                text,
                first_indent=False,
                center=True,
                font_size=BODY_SIZE,
                line_spacing=BODY_LINE_SPACING,
                paragraph=p,
            )
            return
        _set_no_midword_wrap(p)
        _append_text_runs(p, str(text), size=BODY_SIZE, bold=bold)

    def add_keywords(self, text: str):
        p = self.doc.add_paragraph()
        _para_spacing(p.paragraph_format, first_indent=False)
        run = p.add_run("关键词：")
        set_heading_font(run, BODY_SIZE)
        run = p.add_run(text)
        set_run_font(run, BODY_SIZE, False)

    def add_table(self, caption: str, headers: list[str], rows: list[list[str]]):
        self.tbl_no += 1
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        set_run_font(run, CAPTION_SIZE, True)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.autofit = True
        for j, h in enumerate(headers):
            self._fill_cell(table.rows[0].cells[j], str(h), bold=True)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                self._fill_cell(table.rows[i + 1].cells[j], str(val))
        _apply_three_line_table(table)
        self.doc.add_paragraph()

    def add_figure(self, path: Path | str, caption: str, *, width=5.2):
        path = Path(path)
        self.fig_no += 1
        cap_text = f"图 {self.fig_no}  {caption}"
        if not path.is_file():
            self.add_para(f"[缺图: {path.name}]", first_indent=False)
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(cap_text)
            set_run_font(run, CAPTION_SIZE, False)
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(cap_text)
        set_run_font(run, CAPTION_SIZE, False)

    def try_add_figure(self, candidates: list[Path | str], caption: str, *, width=5.2):
        for c in candidates:
            if Path(c).is_file():
                self.add_figure(c, caption, width=width)
                return
        self.add_figure(Path(str(candidates[0])), caption, width=width)

    def add_refs(self, refs: dict[int, str]):
        self.add_heading1("参考文献")
        for n in sorted(refs):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(f"[{n}] {refs[n]}")
            set_run_font(run, REF_SIZE, False)

    def validate_cites(self, refs: dict[int, str]):
        body = self.cites_used
        ref_nums = set(refs)
        if body != ref_nums:
            missing_in_refs = body - ref_nums
            unused_refs = ref_nums - body
            print(f"  cite check: body={sorted(body)} refs={sorted(ref_nums)}")
            if missing_in_refs:
                print(f"  WARNING missing refs for {sorted(missing_in_refs)}")
            if unused_refs:
                print(f"  WARNING unused refs {sorted(unused_refs)}")
